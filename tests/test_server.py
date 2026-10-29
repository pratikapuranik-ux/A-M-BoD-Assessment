"""Server tests: framework base parity with js/data/base.js, views, import end to end, API smoke on an empty start, mocked research.
The Chevron runs in skill/example_runs are used only as test fixtures; the app itself starts with no companies.

    python -m pytest tests/test_server.py      # or: python tests/test_server.py
"""
import base64
import copy
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import types

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "server"))
# Short temp path: the skill scripts write with plain open(), so stay well under Windows MAX_PATH.
TMP = tempfile.mkdtemp(prefix="bd_")
os.environ["BOARD_DATA_DIR"] = TMP
os.environ["BOARD_MIN_RESPONSES"] = "3"  # test the real guardrail, whatever .env says for demos

import convert  # noqa: E402
import extract  # noqa: E402
import jobs  # noqa: E402

EX = os.path.join(ROOT, "skill", "example_runs")
_client = None


def load(name):
    with open(os.path.join(EX, name), encoding="utf-8") as f: return json.load(f)


def js_global(path, name):
    """Evaluate a js/data/*.js file with node and return the named global as JSON."""
    code = f"const fs=require('fs');const D=new Function(fs.readFileSync(process.argv[1],'utf8')+';return {name}')();process.stdout.write(JSON.stringify(D))"
    r = subprocess.run(["node", "-e", code, os.path.join(ROOT, path)], capture_output=True, text=True, encoding="utf-8", check=True)
    return json.loads(r.stdout)


def norm(s):
    return re.sub(r"\s+", " ", s.replace("’", "'").replace("‘", "'")).strip() if isinstance(s, str) else s


def diff(a, b, p="D", strict=False):
    """Differences for every key present in a (the reference); extra keys in b are allowed."""
    out = []
    if isinstance(a, dict):
        if not isinstance(b, dict): return [(p, a, b)]
        for k in a:
            out += [(f"{p}.{k}", a[k], None)] if k not in b else diff(a[k], b[k], f"{p}.{k}", strict)
    elif isinstance(a, list):
        if not isinstance(b, list) or len(a) != len(b): return [(p, a, b)]
        for i, (x, y) in enumerate(zip(a, b)): out += diff(x, y, f"{p}[{i}]", strict)
    elif (a != b) if strict else (norm(a) != norm(b)):
        out.append((p, a, b))
    return out


def client():
    global _client
    if _client is None:
        from fastapi.testclient import TestClient
        import app
        _client = TestClient(app.app); _client.__enter__()
    return _client


def wait_job(c, jid):
    jobs.wait(jid, 600); return c.get(f"/api/jobs/{jid}").json()


# ---------------- convert parity ----------------

def test_base_view_matches_base_js():
    assert js_global("js/data/base.js", "D") == json.loads(json.dumps(convert.base_view())), "js/data/base.js is stale: run python tools/build_base.py"
    B = convert.base_view()
    assert B["meta"]["company"] == "" and B["peers"] == [] and B["io"] is None and len(B["ind"]) == 22 and all(c["score"] is None for c in B["ind"].values())


def test_view_from_fixture_runs():
    V = json.loads(json.dumps(convert.to_view(load("outside_in_run.json"), load("inside_out_run.json"), None, {"id": "fixture"})))
    assert V["meta"]["prefixes"] == {"E": "Chevron", "X": "ExxonMobil", "B": "BP"} and V["meta"]["id"] == "fixture"
    assert set(V) >= set(convert.base_view()) and V["defs"] == convert.base_view()["defs"]
    assert "simulated" not in json.dumps(V["meta"]) and "final" not in V["io"]


def test_framework_matches_framework_js():
    assert convert.framework() == js_global("js/data/framework.js", "FW")


def test_view_without_io_or_prior():
    oi = load("outside_in_run.json"); oi.pop("history", None)
    V = convert.to_view(oi)
    assert V["io"] is None and V["prior"] is None and V["findings_io"] == []
    assert V["meta"]["short"] == "Chevron" and V["meta"]["peer_short"] == ["ExxonMobil", "BP"]


def test_short_names():
    assert convert.short_name("Exxon Mobil (ExxonMobil Holdings Corporation)") == "ExxonMobil"
    assert convert.short_name("BP p.l.c.") == "BP" and convert.short_name("Chevron Corporation") == "Chevron"
    assert convert.short_name("Apple Inc.") == "Apple"


# ---------------- extract ----------------

def test_extract_docx_txt():
    from docx import Document
    doc = Document(); doc.add_paragraph("Minutes of meeting 1"); t = doc.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Item"; t.rows[0].cells[1].text = "Buyback"
    buf = io.BytesIO(); doc.save(buf)
    x = extract.extract_b64({"slot": 1, "name": "m.docx", "data_b64": base64.b64encode(buf.getvalue()).decode()})
    assert x["slot"] == 1 and "Minutes of meeting 1" in x["text"] and "Item | Buyback" in x["text"] and x["pdf_b64"] is None
    assert extract.extract("a.txt", "café".encode(), 0)["text"] == "café"
    try: extract.extract("a.exe", b"x"); raise AssertionError("expected ExtractError")
    except extract.ExtractError: pass


# ---------------- API ----------------

def test_api_smoke():
    c = client()
    h = c.get("/api/health").json(); assert h["ok"] is True and "enabled" in h["research"] and "reason" in h["research"]
    assert c.get("/api/framework").json()["rule_version"] == "1.2"
    assert "seed_job" not in h and c.get("/api/companies").json() == [], "a fresh start must have no companies"
    r = c.get("/api/companies/nope/view"); assert r.status_code == 404 and "error" in r.json()
    assert c.get("/api/jobs/nope").status_code == 404 and c.delete("/api/companies/nope").status_code == 404
    assert c.get("/").status_code == 200 and c.get("/js/engine.js").status_code == 200 and c.get("/js/data/base.js").status_code == 200
    assert c.get("/js/data/chevron.js").status_code == 404 and c.get("/js/data/downloads.js").status_code == 404


def renamed(run, name="Testco Corporation", ticker="TST"):
    run = copy.deepcopy(run); comp = convert.company_of(run); comp["name"] = name; comp["ticker"] = ticker; run.pop("history", None); run["run"].pop("simulated", None); return run


def test_import_end_to_end():
    c = client()
    bad = renamed(load("outside_in_run.json")); convert.company_of(bad)["indicators"]["1.1"]["board"]["score"] = 7
    r = c.post("/api/companies/import", json={"outside_in": bad}); assert r.status_code == 422 and any("not an integer" in p for p in r.json()["problems"])
    sim = renamed(load("outside_in_run.json")); sim["run"]["simulated"] = True
    r = c.post("/api/companies/import", json={"outside_in": sim}); assert r.status_code == 422 and any("simulated" in p for p in r.json()["problems"])
    oi = renamed(load("outside_in_run.json")); prior = renamed(load("prior_2026-Q3_simulated_run.json"))
    r = c.post("/api/companies/import", json={"outside_in": oi, "priors": [prior]}); assert r.status_code == 200, r.text
    body = r.json(); assert body["id"] == "testco"
    j = wait_job(c, body["job_id"]); assert j["status"] == "done", j
    assert c.post("/api/companies/import", json={"outside_in": oi}).status_code == 409
    s = next(x for x in c.get("/api/companies").json() if x["id"] == "testco")
    assert s["stages"] == {"outside_in": True, "inside_out": False} and s["source"] == "imported" and s["prior_period"] == "2026-Q3"
    v = c.get("/api/companies/testco/view").json()
    assert v["meta"]["company"] == "Testco Corporation" and v["meta"]["short"] == "Testco" and v["io"] is None and v["prior"]["period"] == "2026-Q3"
    assert set(v["files"]) == {"oi_docx", "oi_html", "prior_docx", "quest"}
    for k, url in v["files"].items():
        r = c.get(url); assert r.status_code == 200 and len(r.content) > 1000 and "attachment" in r.headers.get("content-disposition", ""), k
    assert c.get("/api/companies/testco/view?mode=internal_dry_run").json()["meta"]["mode"] == "internal_dry_run"
    assert c.get("/api/companies/testco/view?mode=bogus").status_code == 400 and c.get("/api/companies/testco/files/nope").status_code == 404
    assert [(p["period"], p["stage"]) for p in v["priors"]] == [("2026-Q3", "outside_in")]
    import store  # a later quarter for the same company: the 2026-Q4 assessment joins the earlier ones
    later = copy.deepcopy(oi); later["run"]["period"] = "2027-Q1"; store.new_period("testco", later)
    v = c.get("/api/companies/testco/view").json()
    assert v["meta"]["period"] == "2027-Q1" and [p["period"] for p in v["priors"]] == ["2026-Q3", "2026-Q4"] and v["prior"]["period"] == "2026-Q4"
    assert c.delete("/api/companies/testco").json()["ok"] and c.get("/api/companies/testco/view").status_code == 404


def test_research_mocked():
    c = client()
    calls = {}
    def run_outside_in(company, ticker, peers, period, mode, progress):
        calls["oi"] = (company, ticker, peers, period, mode); progress(50, "half way"); return renamed(load("outside_in_run.json"), "Mockco Inc.", "MCK")
    def run_inside_out(oi_run, documents, answers, ratings, progress):
        calls["io"] = (documents, answers, ratings); progress(50, "reading")
        io_run = renamed(load("inside_out_run.json"), "Mockco Inc.", "MCK"); io_run["run"]["simulated"] = False; io_run.pop("self_assessment_answers", None); return io_run
    sys.modules["research"] = types.SimpleNamespace(research_enabled=lambda: (True, "mock"), run_outside_in=run_outside_in, run_inside_out=run_inside_out)
    try:
        assert c.get("/api/health").json()["research"] == {"enabled": True, "reason": "mock"}
        r = c.post("/api/research", json={"company": "Mockco Inc.", "ticker": "MCK", "peers": [{"name": "A"}], "period": "2026-Q4", "mode": "director_request"})
        assert r.status_code == 400 and r.json()["problems"]
        r = c.post("/api/research", json={"company": "Mockco Inc.", "ticker": "MCK", "peers": [{"name": "A", "ticker": "A"}, {"name": "B", "ticker": "B"}], "period": "2026-Q4", "mode": "director_request"})
        j = wait_job(c, r.json()["job_id"]); assert j["status"] == "done", j
        cid = j["result"]["company_id"]; assert cid == "mockco" and calls["oi"][3] == "2026-Q4"
        assert any("half way" in x for x in j["log"])
        assert next(x for x in c.get("/api/companies").json() if x["id"] == cid)["source"] == "research"
        doc = {"slot": 1, "name": "minutes.txt", "mime": "text/plain", "data_b64": base64.b64encode(b"Board minutes: buyback approved").decode()}
        assert c.post(f"/api/companies/{cid}/inside-out", json={"documents": [dict(doc, slot=99)]}).status_code == 400
        # anonymity guardrail: no inside-out view until three people have answered the self-assessment
        r = c.post(f"/api/companies/{cid}/inside-out", json={"mode": "director_request", "documents": [doc]})
        assert r.status_code == 409 and "0 of 3" in r.json()["error"]
        st = c.post(f"/api/companies/{cid}/survey/invite", json={"emails": ["a@example.org", "b@example.org", "c@example.org", "a@example.org"]}).json()
        assert len(st["invites"]) == 3 and st["responses"] == 0 and all("?take=" in i["link"] for i in st["invites"])
        assert c.post(f"/api/companies/{cid}/survey/invite", json={"emails": ["not-an-email"]}).status_code == 400
        tok = [i["token"] for i in st["invites"]]
        form = c.get(f"/api/survey/{tok[0]}").json()
        assert form["company"] == "Mockco Inc." and len(form["questions"]) == 12 and set(form) == {"company", "period", "responded", "min", "questions"}
        for k, t in enumerate(tok[:2]):
            assert c.post(f"/api/survey/{t}", json={"answers": {"1": f"Customer item {k}"}, "ratings": {"1": 3}}).status_code == 200
        assert c.post(f"/api/survey/{tok[0]}", json={"answers": {"1": "again"}}).status_code == 409
        assert c.post(f"/api/companies/{cid}/inside-out", json={"documents": [doc]}).status_code == 409
        assert c.post(f"/api/companies/{cid}/survey/responses", json={"answers": {"2": "Paper return"}, "ratings": {"2": 5}}).json()["responses"] == 3
        stored = json.dumps(json.load(open(os.path.join(TMP, "companies", cid, "survey.json"), encoding="utf-8"))["responses"])
        assert not any(t in stored for t in tok) and "example.org" not in stored  # responses carry no link to who answered
        assert c.get("/api/survey/nope").status_code == 404
        r = c.post(f"/api/companies/{cid}/inside-out", json={"mode": "director_request", "documents": [doc]})
        assert r.status_code == 200, r.text
        j = wait_job(c, r.json()["job_id"]); assert j["status"] == "done", j
        d = calls["io"][0][0]; assert d["slot"] == 1 and "buyback" in d["text"] and d["pdf_b64"] is None
        a1 = calls["io"][1]["1"]; assert "Customer item 0" in a1 and "Customer item 1" in a1 and a1.count("Response ") == 2 and re.fullmatch(r"Response [123]: Paper return", calls["io"][1]["2"])  # responses are shuffled, so the number varies
        v = c.get(f"/api/companies/{cid}/view").json()
        assert "Customer item" in v["io"]["sa"][0] and "io_docx" in v["files"] and "pack" not in v["files"]
        assert c.delete(f"/api/companies/{cid}").status_code == 200
    finally:
        sys.modules.pop("research", None)


# ---------------- fixes: thresholds, ids, periods, jobs, files, survey, uploads, reports ----------------

SCRIPTS = os.path.join(ROOT, "skill", "scripts")


def scripts():
    if SCRIPTS not in sys.path: sys.path.insert(0, SCRIPTS)


def test_two_thirds_threshold_is_exact():
    scripts(); import common, grade
    assert common.meets(2, 3, 0.67) and not common.meets(1, 3, 0.67) and common.meets(2, 4, 0.5) and not common.meets(1, 3, 0.5)
    subj = {"role": "company", "indicators": {i: {"board": {"score": 3}, "board_io": {"score": 3}} for i in common.IND_IDS}}
    d3 = [i for i in common.IND_IDS if common.DIM_OF[i] == "4"]; assert len(d3) == 3
    subj["indicators"][d3[0]]["board_io"] = {"score": None, "na_cause": "retrieval-limited"}
    assert grade.compute_subject(subj, ["board", "board_io"], "inside_out")["dimensions"]["4"]["board_io"]["reported"] is True


def test_company_ids():
    import store
    assert convert.slug("株式会社トヨタ", "7203") == "7203" and re.fullmatch(r"co-[0-9a-f]{10}", convert.slug("Сбербанк"))
    assert convert.slug("Сбербанк") != convert.slug("Газпром")
    long = "Samsung Electronics Co., Ltd. and Very Long Subsidiary Name International Holdings Incorporated"
    s = convert.slug(long); assert len(s) <= 56 and store.valid_id(s) and store.valid_id(s + "-99")
    r = client().post("/api/companies/import", json={"outside_in": renamed(load("outside_in_run.json"), long, "SMSN")})
    assert r.status_code == 200, r.text
    assert wait_job(client(), r.json()["job_id"])["status"] == "done" and store.exists(r.json()["id"])
    assert client().delete(f"/api/companies/{r.json()['id']}").status_code == 200


def test_import_rejects_bad_priors():
    oi = renamed(load("outside_in_run.json"), "Priorco"); p = renamed(load("prior_2026-Q3_simulated_run.json"), "Priorco")
    later = copy.deepcopy(p); later["run"]["period"] = "2027-Q1"
    r = client().post("/api/companies/import", json={"outside_in": oi, "priors": [p, copy.deepcopy(p), later]})
    assert r.status_code == 422 and any("more than once" in x for x in r.json()["problems"]) and any("2027-Q1 is not earlier" in x for x in r.json()["problems"])


def test_new_period_rules():
    import store
    oi = renamed(load("outside_in_run.json"), "Periodco", "PRD"); io = renamed(load("inside_out_run.json"), "Periodco", "PRD")
    cid = store.create(oi, io); store.add_response(cid, {"1": "a"}, {})
    assert store.new_period(cid, copy.deepcopy(oi)) is not None, "a same-period refresh returns a backup"
    assert len(store.load_survey(cid)["responses"]) == 1 and store.read_run(cid, "inside_out") and store.prior_paths(cid) == []
    early = copy.deepcopy(oi); early["run"]["period"] = "2026-Q3"
    try: store.new_period(cid, early); raise AssertionError("expected ValueError")
    except ValueError as e: assert "earlier than the 2026-Q4" in str(e)
    later = copy.deepcopy(oi); later["run"]["period"] = "2027-Q1"
    assert store.new_period(cid, later) is None
    assert not os.path.exists(store.survey_path(cid)) and store.read_run(cid, "inside_out") is None
    assert [store._read(p)["run"]["stage"] for p in store.prior_paths(cid)] == ["inside_out"]
    try: store.new_period(cid, copy.deepcopy(oi)); raise AssertionError("expected ValueError")
    except ValueError: pass
    store.delete(cid)


def mock_research(**kw):
    class ResearchError(RuntimeError): pass
    class ResearchInvalid(ResearchError):
        def __init__(self, msg, problems=()): super().__init__(msg); self.problems = list(problems)
    return types.SimpleNamespace(research_enabled=lambda: (True, "mock"), ResearchError=ResearchError, ResearchInvalid=ResearchInvalid, **kw)


def test_jobs_ids_and_refresh_via_research():
    import threading, store
    c = client(); gate = threading.Event(); mode = {"bad": False, "raise": None}
    def run_outside_in(company, ticker, peers, period, m, progress):
        gate.wait(30)
        if mode["raise"]: raise mode["raise"]
        r = renamed(load("outside_in_run.json"), company, ticker); r["run"]["period"] = period
        if mode["bad"]: convert.company_of(r)["indicators"]["1.1"]["board"]["score"] = 7
        return r
    mod = sys.modules["research"] = mock_research(run_outside_in=run_outside_in, run_inside_out=None)
    peers = [{"name": "A", "ticker": "A"}, {"name": "B", "ticker": "B"}]
    go = lambda name, ticker="GTC", period="2026-Q4": c.post("/api/research", json={"company": name, "ticker": ticker, "peers": peers, "period": period})
    try:
        a = go("Gateco Inc.").json()["job_id"]
        r = go("Gateco Inc."); assert r.status_code == 409 and r.json()["job_id"] == a, "a second start returns the live job"
        gate.set(); j = wait_job(c, a); assert j["status"] == "done" and j["result"]["company_id"] == "gateco", j
        j = wait_job(c, go("Gateco Ltd", "GTX").json()["job_id"]); assert j["result"]["company_id"] == "gateco-2", "a different company with the same slug"
        assert store.load_meta("gateco")["ticker"] == "GTC"
        gate.clear(); a = go("Gateco Inc.").json()["job_id"]
        r = c.delete("/api/companies/gateco"); assert r.status_code == 409 and r.json()["error"] == "a job is running for this company"
        r = c.post("/api/companies/gateco/inside-out", json={}); assert r.status_code == 409 and r.json()["job_id"] == a
        gate.set(); assert wait_job(c, a)["status"] == "done"
        r = go("Gateco Inc.", period="2026-Q3"); assert r.status_code == 409 and "earlier than the 2026-Q4" in r.json()["error"]
        before = store.read_run("gateco", "outside_in"); mode["bad"] = True
        j = wait_job(c, go("Gateco Inc.").json()["job_id"]); assert j["status"] == "error" and "previous 2026-Q4 assessment was kept" in j["error"]
        assert store.read_run("gateco", "outside_in") == before and "problems" not in store.load_meta("gateco")
        mode["bad"] = False; mode["raise"] = mod.ResearchInvalid("The scores still break 1 rule(s).", [("Gateco Inc.", "1.1/board", "score 5 needs more")])
        j = wait_job(c, go("Gateco Inc.").json()["job_id"])
        assert j["error"] == "The scores still break 1 rule(s)." and j["problems"] == ["[Gateco Inc.] 1.1/board: score 5 needs more"]
        for cid in ("gateco", "gateco-2"): assert c.delete(f"/api/companies/{cid}").status_code == 200
    finally:
        gate.set(); sys.modules.pop("research", None)


def test_research_with_own_documents():
    import store
    c = client(); got = {}
    def run_outside_in(company, ticker, peers, period, m, progress, documents=None):
        got["docs"] = documents; r = renamed(load("outside_in_run.json"), company, ticker); r["run"]["period"] = period; return r
    sys.modules["research"] = mock_research(run_outside_in=run_outside_in, run_inside_out=None)
    body = {"company": "Docco plc", "ticker": "DOC", "peers": [{"name": "A"}, {"name": "B"}], "period": "2026-Q4"}
    try:
        r = c.post("/api/research", json=dict(body, documents=[{"name": "notes.txt", "data_b64": base64.b64encode(b"Board met twice.").decode()}]))
        j = wait_job(c, r.json()["job_id"]); assert j["status"] == "done", j
        assert [(d["name"], d["text"]) for d in got["docs"]] == [("notes.txt", "Board met twice.")]
        r = c.post("/api/research", json=dict(body, documents=[{"name": "x.exe", "data_b64": "eA=="}]))
        assert r.status_code == 400 and "x.exe: unsupported file type" in r.json()["problems"][0]
        assert c.delete(f"/api/companies/{j['result']['company_id']}").status_code == 200
    finally:
        sys.modules.pop("research", None)


def test_store_file_locks(monkeypatch):
    import store
    real, n = os.replace, {"k": 0}
    def flaky(a, b):
        n["k"] += 1
        if n["k"] < 3: raise PermissionError("held by OneDrive")
        real(a, b)
    monkeypatch.setattr(store.os, "replace", flaky)
    p = os.path.join(TMP, "w.json"); store._write(p, {"a": 1}); monkeypatch.undo()
    assert store._read(p) == {"a": 1} and n["k"] == 3 and not [f for f in os.listdir(TMP) if f.endswith(".tmp")]
    os.makedirs(os.path.join(store.companies_dir(), "broken"), exist_ok=True)
    with open(os.path.join(store.companies_dir(), "broken", "meta.json"), "w") as f: f.write("{")
    oi = renamed(load("outside_in_run.json"), "Listco"); cid = store.create(oi)
    r = client().get("/api/companies"); ids = [x["id"] for x in r.json()]
    assert r.status_code == 200 and cid in ids and "broken" not in ids, "an unreadable company is skipped, not a 500"
    shutil.rmtree(os.path.join(store.companies_dir(), "broken")); store.delete(cid)


def test_leftovers_are_never_adopted(monkeypatch):
    import store
    d = os.path.join(store.companies_dir(), "leftco"); os.makedirs(d)
    for f, obj in (("survey.json", {"invites": [], "responses": [{"answers": {"1": "old"}, "ratings": {}}]}), ("prior_2026-Q1.json", {"run": {}})):
        with open(os.path.join(d, f), "w") as fh: json.dump(obj, fh)
    monkeypatch.setattr(store, "_rmtree", lambda _d: False)  # OneDrive holds the folder
    cid = store.create(renamed(load("outside_in_run.json"), "Leftco"), cid="leftco")
    assert store.load_survey(cid)["responses"] == [] and store.prior_paths(cid) == []
    assert any(x.startswith("leftco.deleted-") for x in os.listdir(store.companies_dir())) and store.list_ids().count("leftco") == 1
    store.add_response(cid, {"1": "secret"}, {}); store.delete(cid)
    assert not os.path.exists(store.survey_path(cid)) and not store.exists(cid), "answers and meta go even if the folder stays"


def test_survey_guardrails(monkeypatch):
    import store, app
    c = client(); cid = store.create(renamed(load("outside_in_run.json"), "Surveyco plc", "SVY")); u = f"/api/companies/{cid}/survey"
    assert c.post(u + "/responses", json={"answers": {}, "ratings": {}}).status_code == 400
    assert c.post(u + "/responses", json={"answers": {"1": "Paper  one"}, "ratings": {"1": 3}}).status_code == 200
    r = c.post(u + "/responses", json={"answers": {"1": "paper one "}, "ratings": {"1": 3}}); assert r.status_code == 409 and r.json()["error"] == "this questionnaire has already been added"
    assert c.post(u + "/responses", json={"answers": {}, "ratings": {"2": 5}}).status_code == 200, "ratings alone are a return"
    assert c.post(u + "/responses", json={"answers": {"3": "x\x0by\x00"}, "ratings": {"1": True, "2": 3.0}}).status_code == 200
    s = store.load_survey(cid)
    assert {"answers": {**{str(k): "" for k in range(1, 13)}, "3": "xy"}, "ratings": {}} in s["responses"], "control characters dropped, bool and float ratings ignored"
    assert all(set(r) == {"answers", "ratings"} for r in s["responses"]) and len(s["hashes"]) == 3
    # emails
    assert c.post(u + "/invite", json={"emails": "a@x.org"}).status_code == 400 and c.post(u + "/invite", json={"emails": [1]}).status_code == 400
    for bad in ["<>", 'a"b@x.org', "a@x.org.", "x" * 250 + "@x.org", "a<b@x.org"]:
        assert c.post(u + "/invite", json={"emails": [bad]}).status_code == 400, bad
    assert c.post(u + "/invite", json={"emails": [f"p{k}@x.org" for k in range(51)]}).status_code == 400
    st = c.post(u + "/invite", json={"emails": ["Ann Lee <ann@x.org>", "bob@x.org"]}).json()
    assert [i["email"] for i in st["invites"]] == ["ann@x.org", "bob@x.org"] and st["local"] is False
    monkeypatch.setattr(store, "MAX_INVITES", 3); assert c.post(u + "/invite", json={"emails": ["c@x.org", "d@x.org"]}).status_code == 400; monkeypatch.undo()
    monkeypatch.setenv("BOARD_PUBLIC_URL", "http://127.0.0.1:5173/")
    st = c.get(u).json(); assert st["invites"][0]["link"].startswith("http://127.0.0.1:5173/?take=") and st["local"] is True
    monkeypatch.setenv("BOARD_PUBLIC_URL", "https://board.example.org"); st = c.get(u).json()
    assert st["invites"][0]["link"].startswith("https://board.example.org/?take=") and st["local"] is False
    # an online response identical to a stored one is refused and leaves the link usable
    tok = st["invites"][0]["token"]
    r = c.post(f"/api/survey/{tok}", json={"answers": {"1": "Paper one"}, "ratings": {"1": 3}}); assert r.status_code == 409 and "identical" in r.json()["error"]
    assert c.post(f"/api/survey/{tok}", json={"answers": {"1": "Online one"}}).status_code == 200
    # differencing: a refresh needs the same responses or MIN_RESPONSES more
    st = c.get(u).json(); assert (st["count"], st["min"], st["can_generate"], st["reason"]) == (4, 3, True, None)
    store.set_used(cid, 4); assert c.get(u).json()["can_generate"] is True
    assert c.post(u + "/responses", json={"answers": {"1": "Fifth"}}).status_code == 200
    st = c.get(u).json(); assert st["can_generate"] is False and st["reason"].startswith("Wait for 2 more responses")
    r = c.post(f"/api/companies/{cid}/inside-out", json={"documents": []}); assert r.status_code == 409 and r.json()["error"] == st["reason"]
    # each question shuffled and numbered on its own
    a, _ = app._combined([{"answers": {"1": "a", "2": ""}, "ratings": {}}, {"answers": {"1": "", "2": "b"}, "ratings": {}}])
    assert a["1"] == "Response 1: a" and a["2"] == "Response 1: b"
    store.delete(cid)


def test_upload_limits(monkeypatch):
    import store
    from pypdf import PdfWriter
    from docx import Document
    c = client(); cid = store.create(renamed(load("outside_in_run.json"), "Uploadco")); u = f"/api/companies/{cid}/inside-out"
    b64 = lambda b: base64.b64encode(b).decode()
    w = PdfWriter()
    for _ in range(101): w.add_blank_page(72, 72)
    buf = io.BytesIO(); w.write(buf)
    r = c.post(u, json={"documents": [{"slot": 1, "name": "pack.pdf", "data_b64": b64(buf.getvalue())}]})
    assert r.status_code == 400 and any(p.startswith("pack.pdf: a scanned PDF (101 pages)") for p in r.json()["problems"])
    buf = io.BytesIO(); Document().save(buf)
    r = c.post(u, json={"documents": [{"slot": 1, "name": "empty.docx", "data_b64": b64(buf.getvalue())}]})
    assert r.status_code == 400 and "no readable text in empty.docx; export it as PDF or text" in r.json()["problems"]
    monkeypatch.setattr(extract, "MAX_FILE", 10)
    r = c.post(u, json={"documents": [{"slot": 1, "name": "big.txt", "data_b64": b64(b"x" * 11)}]})
    assert r.status_code == 400 and "big.txt: 0 MB; the limit is 0 MB per file" in r.json()["problems"]
    monkeypatch.setattr(extract, "MAX_FILE", 25 * 2**20)
    monkeypatch.setattr(extract, "MAX_UPLOAD", 30)
    r = c.post(u, json={"documents": [{"slot": 1, "name": "a.txt", "data_b64": b64(b"x" * 20)}, {"slot": 2, "name": "b.txt", "data_b64": b64(b"y" * 20)}]})
    assert r.status_code == 400 and r.json()["error"].startswith("b.txt: the documents together exceed")
    store.delete(cid)


def test_fit_pdfs_sends_overflow_as_text(monkeypatch):
    monkeypatch.setattr(extract, "PDF_INLINE_PAGES", 100)
    pdf = lambda n, p, size, text="[page 1]\nminutes": {"name": n, "pages": p, "size": size, "text": text, "pdf_b64": "x"}
    docs = extract.fit_pdfs([pdf("pack.pdf", 300, 20 * 2**20), pdf("minutes.pdf", 40, 2**20), pdf("papers.pdf", 70, 2**20)])
    assert [d["pdf_b64"] is not None for d in docs] == [False, True, False]  # smallest first, within 100 pages
    with pytest.raises(extract.ExtractError, match="scanned PDF"):
        extract.fit_pdfs([pdf("scan.pdf", 150, 5 * 2**20, "[page 1]\n\n[page 2]")])


def test_report_helpers():
    scripts(); import docx_helpers as H, build_docx
    from docx import Document
    d = Document(); H.para(d, "answer\x0bline"); H.set_cell(d.add_table(rows=1, cols=1).rows[0].cells[0], "a\x01b")
    p = d.add_paragraph(); H.text_with_links(p, "peers [X01] and [K12], company [E03], board [IO01]\x0c")
    xml = p._p.xml; assert all(f'w:anchor="EV_{c}"' in xml for c in ("X01", "K12", "E03", "IO01"))
    assert "[EXBKLN]" in open(os.path.join(SCRIPTS, "build_dashboard.py"), encoding="utf-8").read()
    h = {"stage": "inside_out", "dimensions": {"1": {"board": {"grade": "D"}, "board_io": {"grade": "B"}}}, "overall": {"board": {"grade": "D"}, "board_io": {"grade": "C"}}}
    assert build_docx.hist_grade(h, "1") == "B" and build_docx.hist_grade(h) == "C" and build_docx.hist_grade(dict(h, stage="outside_in"), "1") == "D"


def teardown_module(_m=None):
    global _client
    if _client is not None: _client.__exit__(None, None, None); _client = None
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    order = ["test_short_names", "test_framework_matches_framework_js", "test_base_view_matches_base_js", "test_view_from_fixture_runs", "test_view_without_io_or_prior", "test_extract_docx_txt", "test_api_smoke", "test_import_end_to_end", "test_research_mocked"]
    failed = 0
    for n in order:
        try: globals()[n](); print("PASS", n)
        except Exception as e:  # noqa: BLE001
            failed += 1; import traceback; print("FAIL", n); traceback.print_exc()
    teardown_module()
    print(f"{len(order) - failed}/{len(order)} passed"); sys.exit(1 if failed else 0)
