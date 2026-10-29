"""In-process deliverable generators (no subprocesses): Word report, HTML dashboard, questionnaire.

Ported from skill/scripts (build_docx.py, build_dashboard.py, build_questionnaire.py, parse_questionnaire.py,
docx_helpers.py, common.py). Generated artefacts keep the legacy layout and the DRAFT disclaimers.
"""
from __future__ import annotations

import copy
import html
import json
import math
import os
import random
import re
from fractions import Fraction
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- framework / common
_HERE = Path(__file__).resolve().parent
_FW_CANDIDATES = [os.environ.get("AM_FRAMEWORK_JSON") or "", _HERE.parent / "data" / "framework.json",
                  *(p / "skill" / "assets" / "framework.json" for p in _HERE.parents)]


def _load_fw() -> dict:
    for c in _FW_CANDIDATES:
        if c and Path(c).is_file() and Path(c).stat().st_size > 0:
            with open(c, encoding="utf-8") as f: return json.load(f)
    raise FileNotFoundError("framework.json not found (looked in app/data and skill/assets)")


FW = _load_fw()
DIMS = sorted(FW["dimensions"].keys(), key=int)
IND = FW["indicators"]; IND_IDS = [i["id"] for i in IND]
DIM_OF = {i["id"]: i["dimension"] for i in IND}
LV = FW["letter_values"]; VL = {v: k for k, v in LV.items()}
NAVY, BLUE, AMBER, TEAL, SLATE, GREY = "002B49", "0085CA", "CF7F00", "00677F", "5E8AB4", "646464"
TINT, TINT2, LIGHT = "EEF7FD", "DAF2FF", "F2F2F2"
RED = "8B2F0B"  # F and score 1: the dark brick red of the reference report
GRADE_FILL = {"A": TEAL, "B": BLUE, "C": SLATE, "D": AMBER, "F": RED, None: "969696"}
SCORE_FILL = {1: RED, 2: AMBER, 3: SLATE, 4: BLUE, 5: TEAL}
DRAFT = "DRAFT ONLY — Requires review and approval by a qualified Alvarez and Marsal professional before use."
# The director's own rating per dimension, on the same 1-5 scale as the indicator scores (index + 1 is the value)
RATING_LABELS = ("1  Day 2", "2  Mostly Day 2", "3  Mixed", "4  Mostly Day 1", "5  Day 1")


def stage_of(run): return run.get("run", {}).get("stage", "outside_in")


def lenses_for(run):
    st = FW["stages"][stage_of(run)]
    return st["lenses"], st["labels"]


def company_of(run): return next(s for s in run["subjects"] if s.get("role") == "company")


def peers_of(run): return [s for s in run["subjects"] if s.get("role") == "peer"]


def period_label(run): return run.get("run", {}).get("period", run.get("run", {}).get("run_date", ""))


def grade_delta(new, old):
    if not new or not old: return None
    return LV[new] - LV[old]


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _deep_clean(o):
    """python-docx rejects XML-incompatible control characters (e.g. a pasted vertical tab); strip them from every string."""
    if isinstance(o, str): return _CTRL.sub("", o)
    if isinstance(o, list): return [_deep_clean(x) for x in o]
    if isinstance(o, dict): return {k: _deep_clean(v) for k, v in o.items()}
    return o


def _prep(run: dict) -> dict:
    """Copy of the run with control characters removed; grades it in-process if `computed` is missing."""
    run = _deep_clean(copy.deepcopy(run))
    if "computed" not in company_of(run):
        try:
            from app.domain import engine
            res = engine.grade(run)
            if isinstance(res, dict) and "subjects" in res: run = res
        except Exception:  # noqa: BLE001
            pass
    if "computed" not in company_of(run): raise ValueError("run has not been graded (company.computed is missing)")
    return run


# ---------------------------------------------------------------- Word helpers (A&M style)
EID_RE = re.compile(r"\[((?:[EXBKLN]|SA|IO)\d{2,3})\]")  # the company (E), peers (X B K L N), self-assessment, board documents


def clean(text): return _CTRL.sub("", text) if isinstance(text, str) else text


def rgb(h): return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)

def borders(table, kind="light", color="D9DEE5"):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        if kind == "none": e.set(qn("w:val"), "nil")
        else: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)

def fix_widths(t, widths_cm):
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm): gc.set(qn("w:w"), str(int(w * 567)))
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None: tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    for row in t.rows:
        for i, w in enumerate(widths_cm): row.cells[i].width = Cm(w)
    t._widths = widths_cm

def new_row(t):
    r = t.add_row(); r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, w in enumerate(getattr(t, "_widths", [])): r.cells[i].width = Cm(w)
    for c in r.cells: cell_margins(c)
    return r.cells

def run_fmt(r, size=9, bold=False, italic=False, color=None, caps=False, spacing=None):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"
    if color: r.font.color.rgb = rgb(color)
    if caps: r.font.all_caps = True
    if spacing is not None:
        rPr = r._r.get_or_add_rPr(); sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), str(spacing)); rPr.append(sp)

_KNOWN_EV: set | None = None  # evidence codes that have an entry in the sources appendix (set per build_docx call); None = link every code


def _linkable(code): return _KNOWN_EV is None or code in _KNOWN_EV


def rich(p, text, size=9.5, bold=False, italic=False, color=None, caps=False, spacing=None):
    """Add text to a paragraph, turning every evidence code that has a source entry into an internal hyperlink."""
    text = text or ""; pos = 0
    for m in EID_RE.finditer(text):
        if not _linkable(m.group(1)): continue
        if m.start() > pos: run_fmt(p.add_run(text[pos:m.start()]), size, bold, italic, color, caps, spacing)
        add_internal_link(p, f"[{m.group(1)}]", f"EV_{m.group(1)}", size, bold=bold); pos = m.end()
    if text[pos:]: run_fmt(p.add_run(text[pos:]), size, bold, italic, color, caps, spacing)


def para(host, text="", size=9.5, bold=False, italic=False, color=None, align=None, after=4, before=0, caps=False, spacing=None, keep=False):
    p = host.add_paragraph()
    if text: rich(p, text, size, bold, italic, color, caps, spacing)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if keep: p.paragraph_format.keep_with_next = True
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def set_cell(cell, text, size=8, bold=False, color=None, align=None, italic=False):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    rich(p, "" if text is None else clean(str(text)), size, bold, italic, color)
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def kicker(doc, text): return para(doc, text, size=7.5, bold=True, color=BLUE, caps=True, spacing=40, after=1, before=8, keep=True)
def headline(doc, text, size=13): return para(doc, text, size=size, bold=True, color=NAVY, after=4, keep=True)
def section(doc, kick, head, size=13): kicker(doc, kick); headline(doc, head, size)

def grade_cell(cell, grade, size=10):
    set_cell(cell, grade or "n/a", size=size, bold=True, color="FFFFFF", align="center"); shade(cell, GRADE_FILL.get(grade))

def score_cell(cell, score, size=9):
    if score is None: set_cell(cell, "n/a", size=7, color=GREY, align="center"); shade(cell, LIGHT)
    else: set_cell(cell, score, size=size, bold=True, color="FFFFFF", align="center"); shade(cell, SCORE_FILL[score])

def table(doc, headers, widths, hsize=7.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, size=hsize, bold=True, color="FFFFFF"); shade(t.rows[0].cells[i], NAVY); cell_margins(t.rows[0].cells[i])
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); trPr.append(th)  # repeat the header row on every page the table spans
    fix_widths(t, widths); return t

def add_bookmark(paragraph, name, bid):
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid)); paragraph._p.insert(0, s); paragraph._p.append(e)

def _link_run(text, size, bold=False):
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    if bold: rPr.append(OxmlElement("w:b"))
    c = OxmlElement("w:color"); c.set(qn("w:val"), BLUE); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial"); rPr.append(rf)
    r.append(rPr); t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = clean(text); r.append(t); return r

def add_internal_link(paragraph, text, anchor, size=7.5, bold=False):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1"); h.append(_link_run(text, size, bold)); paragraph._p.append(h)

def add_external_link(paragraph, url, part, size=7, text=None, bold=False):
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid); h.append(_link_run(text or url, size, bold)); paragraph._p.append(h)

def text_with_links(paragraph, text, evidence_ids=None, size=7.5, lead=None, lead_color=NAVY):
    if lead: run_fmt(paragraph.add_run(lead), size, bold=True, color=lead_color)
    pos = 0; seen = set(); text = text or ""
    for m in EID_RE.finditer(text):
        if not _linkable(m.group(1)): continue
        if m.start() > pos: run_fmt(paragraph.add_run(text[pos:m.start()]), size)
        add_internal_link(paragraph, f"[{m.group(1)}]", f"EV_{m.group(1)}", size); seen.add(m.group(1)); pos = m.end()
    if text[pos:]: run_fmt(paragraph.add_run(text[pos:]), size)
    rest = [e for e in (evidence_ids or []) if e not in seen]
    if rest:
        run_fmt(paragraph.add_run(" Sources: "), size, italic=True, color=GREY)
        for k, e in enumerate(rest):
            if k: run_fmt(paragraph.add_run(", "), size)
            if _linkable(e): add_internal_link(paragraph, f"[{e}]", f"EV_{e}", size)
            else: run_fmt(paragraph.add_run(f"[{e}]"), size)
    paragraph.paragraph_format.space_after = Pt(1)

def add_page_number(paragraph):
    r = paragraph.add_run(); run_fmt(r, 8, color=GREY)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag: f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tag); r._r.append(f)
        else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)

def first_sentence(text):
    text = (text or "").strip(); f, _, _ = text.partition(". "); return (f.rstrip(".") + ".") if f else ""


IND_NAME = {i["id"]: i["name"] for i in IND}


def is_sa(e): return e.get("evidence_class") == "self_assessment" or e.get("source_origin") == "self_assessment"


def mv(dl):
    """Movement as the arrow plus the number of steps (the body shows this instead of an outside-in / inside-out pair)."""
    return "–" if dl is None else "=" if not dl else ("▲ " if dl > 0 else "▼ ") + str(abs(dl))


def mv_color(dl): return TEAL if (dl or 0) > 0 else "A6192E" if (dl or 0) < 0 else GREY


def two_lines(text):
    """First two sentences of existing run prose; nothing is composed here."""
    return " ".join(re.split(r"(?<=[.!?])\s+", (text or "").strip())[:2]) if text else ""


def hist_grade(h, d=None):
    """A history entry's dimension (or overall) grade: board_io for an inside-out quarter, as the dashboard does."""
    x = h["dimensions"][d] if d else h["overall"]
    return (x.get("board_io") or x.get("board") or {}).get("grade") if h["stage"] == "inside_out" else (x.get("board") or {}).get("grade")


# ---------------------------------------------------------------- Word report
def build_docx(run: dict, out_path: str) -> str:
    """Render the Word report from a graded run dict; returns out_path."""
    global _KNOWN_EV
    run = _prep(run); out = out_path; _KNOWN_EV = set(run.get("evidence") or {})
    comp = company_of(run); peers = peers_of(run); C = comp["computed"]
    stage = stage_of(run); lenses, labels = lenses_for(run); io = stage == "inside_out"
    period = period_label(run); hist = run.get("history", []); mode_label = FW["run_modes"].get(run["run"].get("mode"), "")
    simulated_note = " SIMULATED" if run["run"].get("simulated") else ""
    title = "Day 1 / Day 2 Board Assessment" + (": inside-out update" if io else "")

    doc = Document(); st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9.5); st.font.color.rgb = rgb("1A1F2B")
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.6)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; hp.text = ""; run_fmt(hp.add_run("ALVAREZ & MARSAL"), 7, bold=True, color=NAVY, spacing=30)
    run_fmt(hp.add_run(f"   |   {title}   |   {comp['name']}   |   {period}   |   {mode_label}{simulated_note}"), 7, color=GREY)
    fp = sec.footer.paragraphs[0]; fp.text = ""; run_fmt(fp.add_run("▲ " + DRAFT + "      "), 7, italic=True, color=AMBER); add_page_number(fp)

    # ================= COVER =================
    t = doc.add_table(rows=1, cols=1); borders(t, "none"); fix_widths(t, [17.4]); c = t.rows[0].cells[0]; shade(c, NAVY); cell_margins(c, 700, 700, 600, 600); c.text = ""
    r0 = c.paragraphs[0]; run_fmt(r0.add_run("ALVAREZ & MARSAL"), 9, bold=True, color="A8CFEA", spacing=60); r0.paragraph_format.space_after = Pt(60)
    para(c, "Day 1 / Day 2", size=30, bold=True, color="FFFFFF", after=0); para(c, "Board Assessment", size=30, bold=True, color="FFFFFF", after=6)
    if io: para(c, "Inside-out update", size=16, bold=True, color="F2B94B", after=10)
    para(c, f"{comp['name']}" + (f" ({comp.get('ticker')})" if comp.get("ticker") else "") + f"  ·  {period}", size=12, color="FFFFFF", after=4)
    para(c, ("Outside-in assessment from public information" if not io else "Outside-in assessment updated with board documents and the director self-assessment") + f", benchmarked against {', '.join(p['name'] for p in peers)}", size=10, color="BBD6EA", after=30)
    bar = c.add_paragraph(); run_fmt(bar.add_run("────────────"), 10, color="C9A24A"); para(c, run["run"].get("run_date", ""), size=10, bold=True, color="FFFFFF", after=0)
    doc.add_paragraph()
    info = doc.add_table(rows=0, cols=2); borders(info, "none"); fix_widths(info, [3.2, 14.2])
    for k, v in (("Purpose", f"Grade the board of {comp['name']} A-F on six Day 1 / Day 2 dimensions and 22 indicators, attributing what the company did under the board's watch to the board, against a named peer set."),
                 ("Basis", ("Public information only." if not io else "Public information plus the board documents and self-assessment supplied by the director; peers remain on public information.") + f" Rules v{FW['rule_version']}. {mode_label}." + (" SIMULATED RUN." if run["run"].get("simulated") else "")),
                 *([("Trend", f"Compared with {len(hist)} prior quarter(s): {', '.join(h['period'] + (' (simulated)' if h['simulated'] else '') for h in hist)}.")] if hist else []),  # no row without a prior report
                 ("Evidence window", f"Behaviour indicators from {run['run'].get('document_window', {}).get('behaviour_from', '')}; trend indicators from {run['run'].get('document_window', {}).get('trend_from', '')}.")):
        row = new_row(info); set_cell(row[0], k, size=8, bold=True, color=NAVY); shade(row[0], TINT); set_cell(row[1], v, size=8.5)
    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1); borders(box, "none"); fix_widths(box, [17.4]); bc = box.rows[0].cells[0]; shade(bc, "FFF5E6"); cell_margins(bc, 120, 120, 200, 200); bc.text = ""
    run_fmt(bc.paragraphs[0].add_run("DRAFT ONLY — " + mode_label.upper()), 8, bold=True, color=AMBER)
    para(bc, "Grades apply to the board as a body; no individual is assessed. Where a grade rests on outside-in information, additional documentation can arrive at a more accurate grade. Requires review and approval by a qualified Alvarez and Marsal professional before use.", size=8, color="4A4A4A", after=0)
    doc.add_page_break()

    # ================= EXECUTIVE SUMMARY =================
    section(doc, "Executive summary", run.get("headline") or f"{comp['name']}: the board's Day 1 / Day 2 grade", 14)
    shown = lenses[-1:] if io else lenses  # inside-out: one current grade with its movement, not the outside-in / inside-out pair
    tiles = doc.add_table(rows=2, cols=len(shown)); borders(tiles, "none"); fix_widths(tiles, [17.4 / len(shown)] * len(shown))
    for k, l in enumerate(shown):
        o = C["overall"][l]; cell = tiles.rows[0].cells[k]; shade(cell, TINT); cell_margins(cell, 140, 140, 200, 200); cell.text = ""
        run_fmt(cell.paragraphs[0].add_run((labels[l] if not io else "Board grade").upper()), 7.5, bold=True, color=GREY, spacing=30)
        p = cell.add_paragraph(); run_fmt(p.add_run(o.get("grade") or "—"), 30, bold=True, color=GRADE_FILL.get(o.get("grade"))); p.paragraph_format.space_after = Pt(0)
        run_fmt(p.add_run("   " + o.get("label", "")), 9.5, bold=True, color=NAVY)
        if io:
            odl = grade_delta(o.get("grade"), C["overall"]["board"].get("grade")); run_fmt(p.add_run("   " + mv(odl)), 14, bold=True, color=mv_color(odl))
        if o.get("provisional"): run_fmt(p.add_run("   PROVISIONAL"), 7.5, bold=True, color=AMBER)
        prev = next((h for h in reversed(hist) if l in h.get("overall", {}) or l == "board"), None)
        pv = prev["overall"].get(l if l in prev["overall"] else "board", {}).get("grade") if prev else None
        dl = grade_delta(o.get("grade"), pv)
        para(cell, f"{o['dimensions_reported']} of {len(DIMS)} dimensions reported" + (f"  ·  prior quarter {pv} ({'+' if dl > 0 else ''}{dl} step{'s' if abs(dl) != 1 else ''})" if pv and dl is not None else ""), size=7.5, color=GREY, after=0)
        c2 = tiles.rows[1].cells[k]; c2.text = ""; cell_margins(c2, 60, 60, 200, 200); rich(c2.paragraphs[0], two_lines(run.get("what_moved")) if io else (run.get("overall_summary") or {}).get(l, ""), 8.5)  # inside-out: only what changed, from the run's own text
    kicker(doc, "What moved the rating" if (hist or io) else "What drives the rating"); para(doc, run.get("what_moved") or " ".join(run.get("headline_findings", [])), size=9)
    if io and run.get("inside_vs_outside"):
        ivo = run["inside_vs_outside"]; kicker(doc, "Inside view against the public record")
        t2 = table(doc, ["Better than the public evidence suggested", "Worse than the public evidence suggested"], [8.7, 8.7])
        row = new_row(t2)
        for k, key in enumerate(("better", "worse")):
            row[k].text = ""; first = True
            for item in ivo.get(key, []):
                p = row[k].paragraphs[0] if first else row[k].add_paragraph(); first = False; text_with_links(p, "• " + item, size=7.5)
    kicker(doc, "Key actions"); 
    for a in run.get("actions", [])[:5]:
        p = doc.add_paragraph(); rich(p, "▸ " + a["heading"] + ". ", 8.5, bold=True, color=NAVY); rich(p, first_sentence(a["description"]), 8.5); p.paragraph_format.space_after = Pt(1)
    kicker(doc, "Position relative to peers"); para(doc, run.get("peer_position", ""), size=9)
    for ln in run.get("peer_summary_lines", [])[:5]:
        p = doc.add_paragraph(); rich(p, "• " + ln, 8.5); p.paragraph_format.space_after = Pt(1)
    doc.add_page_break()

    # ================= SCORECARD AND TREND =================
    chg = bool(hist) or io  # the Change column and the trend wording only appear when there is something to compare with
    section(doc, "Scorecard and trend" if chg else "Scorecard", "Grades by dimension" + (f", {len(hist)} prior quarter(s)" if hist else ""))
    hcols = [h["period"] for h in hist]
    hdr = ["Dimension"] + [f"{h['period']}{' (sim)' if h['simulated'] else ''}" for h in hist] + [period] + (["Change"] if chg else []) + ["Evidence in this run"]
    wfix = 3.6 + 1.1 * len(hist) + 1.3 + (1.2 if chg else 0); widths = [3.6] + [1.1] * len(hist) + [1.3] + ([1.2] if chg else []) + [17.4 - wfix]
    tt = table(doc, hdr, widths)
    for d in DIMS:
        row = new_row(tt); set_cell(row[0], f"{d}. {FW['dimensions'][d]}", bold=True, color=NAVY); k = 1
        for h in hist:
            grade_cell(row[k], hist_grade(h, d), 8.5); k += 1
        cur = C["dimensions"][d][lenses[-1]]["grade"]; grade_cell(row[k], cur, 9.5); k += 1  # one current grade; the per-lens pair is in the appendix
        dl = grade_delta(cur, hist_grade(hist[-1], d) if hist else None)  # an inside-out prior counts at its board_io grade
        if dl is None and io: dl = grade_delta(cur, C["dimensions"][d]["board"]["grade"])  # no prior quarter: movement from the outside-in grade
        if chg: set_cell(row[k], mv(dl), size=8, align="center", color=mv_color(dl)); k += 1
        set_cell(row[k], (run.get("scorecard_headline") or {}).get(d, ""), size=7.5)
    row = new_row(tt); set_cell(row[0], "Overall", bold=True, color="FFFFFF"); shade(row[0], NAVY); k = 1
    for h in hist: grade_cell(row[k], hist_grade(h), 8.5); k += 1
    grade_cell(row[k], C["overall"][lenses[-1]].get("grade"), 9.5); k += 1
    if chg:
        odl = grade_delta(C["overall"][lenses[-1]].get("grade"), hist_grade(hist[-1]) if hist else None)
        if odl is None and io: odl = grade_delta(C["overall"][lenses[-1]].get("grade"), C["overall"]["board"].get("grade"))
        set_cell(row[k], mv(odl), size=8, align="center", color=mv_color(odl)); k += 1
    set_cell(row[k], "Rounded average of dimension letters (midpoints round down)." + (" * provisional" if C["overall"][lenses[-1]].get("provisional") else ""), size=7, italic=True)
    cv = C["coverage"][lenses[-1]]
    split = ", ".join(f"{n} {c}" for c, n in (cv.get("not_assessable") or {}).items() if n)
    cov = f"Coverage: {cv['assessed']} of {cv['of']} indicators assessed" + (f" (not assessable: {split})" if split else "") + ". "
    para(doc, cov + ("Grades rest on outside-in information; where an indicator is not assessable or scored on thin public evidence, board documents can arrive at a more accurate grade." if not io else "Grades use the board documents and self-assessment supplied. Change is the movement in steps from the prior quarter, or from the outside-in grade where there is none; both lens grades are in the appendix."), size=7.5, italic=True, color=GREY, before=2)
    # what changed table
    changes = []
    if hist:
        prev = hist[-1]
        for i in IND_IDS:
            cur = ((comp["indicators"].get(i) or {}).get("board") or {}).get("score"); pv = (prev["indicators"].get(i) or {}).get("board")
            if cur is not None and pv is not None and cur != pv: changes.append((i, pv, cur))
    if changes:
        kicker(doc, f"What changed since {hist[-1]['period']}" + (" (simulated prior)" if hist[-1]["simulated"] else ""))
        tc = table(doc, ["Indicator", "Prior", "Now", "Δ", "Evidence that moved the score"], [4.6, 1.1, 1.1, 0.9, 9.7])
        notes = run.get("change_notes") or {}
        for i, pv, cur in changes:
            ind = next(x for x in IND if x["id"] == i); row = new_row(tc); set_cell(row[0], f"{i} {ind['name']}", bold=True, color=NAVY)
            score_cell(row[1], pv, 8); score_cell(row[2], cur, 8); set_cell(row[3], f"{'+' if cur > pv else ''}{cur - pv}", align="center", bold=True, color=TEAL if cur > pv else "A6192E")
            p = row[4].paragraphs[0]; text_with_links(p, notes.get(i) or first_sentence(((comp["indicators"][i].get("board") or {}).get("rationale"))), size=7.5)

    # ================= PEER COMPARISON =================
    section(doc, "Peer comparison", "Dimension grades against the peer set" + (" (outside-in grades only)" if io else ""))
    if io: para(doc, "Peers are compared on outside-in grades only; their inside-out information is not available. The company's outside-in grade is used in this table.", size=7.5, italic=True, color=GREY)
    hdr = ["Dimension", comp["name"]] + [p_["name"] for p_ in peers] + ["How the company compares"]
    pw = min(1.5, 7.0 / (1 + len(peers))); widths = [2.9] + [pw] * (1 + len(peers)) + [17.4 - 2.9 - pw * (1 + len(peers))]
    tp = table(doc, hdr, widths, hsize=7); rs = run.get("peer_row_summary") or {}
    for d in DIMS:
        row = new_row(tp); set_cell(row[0], f"{d}. {FW['dimensions'][d]}", bold=True, color=NAVY, size=7.5); grade_cell(row[1], C["dimensions"][d]["board"]["grade"], 9)
        for k, p_ in enumerate(peers, 2): grade_cell(row[k], p_["computed"]["dimensions"][d]["board"]["grade"], 9)
        set_cell(row[len(peers) + 2], rs.get(d, (run.get("comparative_evidence") or {}).get(d, "")), size=7.5)
    row = new_row(tp); set_cell(row[0], "Overall", bold=True, color="FFFFFF", size=7.5); shade(row[0], NAVY); grade_cell(row[1], C["overall"]["board"].get("grade"), 9)
    for k, p_ in enumerate(peers, 2): grade_cell(row[k], p_["computed"]["overall"]["board"].get("grade"), 9)
    set_cell(row[len(peers) + 2], run.get("peer_position", ""), size=7.5)
    if hist and hist[-1].get("peers"):
        ph = hist[-1]["peers"]; kicker(doc, f"Peer trend: {hist[-1]['period']} to {period} (overall)")
        tr = table(doc, ["Company"] + [p_["name"] for p_ in peers], [4.0] + [(13.4) / len(peers)] * len(peers), hsize=7)
        row = new_row(tr); set_cell(row[0], hist[-1]["period"], bold=True)
        for k, p_ in enumerate(peers, 1): grade_cell(row[k], (ph.get(p_["name"]) or {}).get("overall"), 8.5)
        row = new_row(tr); set_cell(row[0], period, bold=True)
        for k, p_ in enumerate(peers, 1): grade_cell(row[k], p_["computed"]["overall"]["board"].get("grade"), 8.5)

    # ================= RECONCILIATION =================
    rec = run.get("reconciliation")
    if rec:
        section(doc, "Reconciliation", "Where the two runs differed")
        mA = (rec.get("models") or {}).get("A") or {}; mB = (rec.get("models") or {}).get("B") or {}
        rows_ = rec.get("rows") or []
        para(doc, f"Run A ({mA.get('model', 'n/a')}, {mA.get('effort', 'n/a')} effort) and Run B ({mB.get('model', 'n/a')}, {mB.get('effort', 'n/a')} effort) agreed on {rec.get('agreed')} of {rec.get('total')} indicator scores; {len(rows_)} were reconciled and the model's recommendation is the final score.", size=8.5, keep=bool(rows_))
        if not rows_:
            para(doc, "Both runs agreed on every indicator score; nothing needed reconciling.", size=8.5, italic=True, color=GREY)
        else:
            tr_ = table(doc, ["Cell", "Score (Run A / Run B)", "Final score", "Reason"], [4.6, 2.4, 2.2, 8.2])
            sc = lambda v: "n/a" if v is None else str(v)
            for r_ in rows_:
                row = new_row(tr_); set_cell(row[0], f"{r_.get('indicator', '')} {r_.get('name', '')}".strip(), bold=True, color=NAVY)
                if r_.get("role") == "peer" and r_.get("subject"): para(row[0], str(r_["subject"]), size=6.5, color=GREY, after=0)
                set_cell(row[1], f"{sc(r_.get('a'))} / {sc(r_.get('b'))}", size=8.5, align="center")
                score_cell(row[2], r_.get("final"), 9)
                if r_.get("chosen") in ("A", "B"): para(row[2], f"Run {r_['chosen']}'s score stands", size=6.5, color="FFFFFF" if r_.get("final") is not None else GREY, after=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell(row[3], r_.get("reason", ""), size=7.5)

    # ================= ACTIONS =================
    section(doc, "Actions", "What the board could do next" + (" (updated after the inside-out review)" if io else ""))
    for k, a in enumerate(run.get("actions", [])[:5], 1):
        p = doc.add_paragraph(); run_fmt(p.add_run(f"{k:02d}  "), 10, bold=True, color=BLUE); rich(p, a["heading"], 9.5, bold=True, color=NAVY); p.paragraph_format.space_after = Pt(1); p.paragraph_format.keep_with_next = True
        p = doc.add_paragraph(); text_with_links(p, a["description"], size=8.5); p.paragraph_format.space_after = Pt(1)
        if a.get("measures") or a.get("target_by"):
            para(doc, ("Measures: " + "; ".join(a.get("measures", [])) if a.get("measures") else "") + (f"   Target: {a['target_by']}" if a.get("target_by") else ""), size=7.5, color=GREY, after=5)
    pas = run.get("prior_actions_status") or []
    if pas:
        kicker(doc, f"Status of actions from {hist[-1]['period'] if hist else 'the prior report'}")
        ts = table(doc, ["Action", "Status", "Evidence and metrics"], [5.6, 2.2, 9.6])
        for pa in pas:
            row = new_row(ts); set_cell(row[0], pa.get("heading", pa.get("id")), bold=True, color=NAVY)
            stt = pa.get("status", ""); col = {"completed": TEAL, "on_track": BLUE, "behind": AMBER, "not_started": "A6192E", "cancelled": GREY, "no_public_evidence": SLATE}.get(stt, GREY)
            set_cell(row[1], stt.replace("_", " "), bold=True, color="FFFFFF", align="center"); shade(row[1], col)
            p = row[2].paragraphs[0]; text_with_links(p, (pa.get("evidence_note") or "") + ((" Metrics: " + pa["metrics_note"]) if pa.get("metrics_note") else ""), size=7.5)
    elif hist and hist[-1].get("actions") and not io:
        kicker(doc, f"Actions from {hist[-1]['period']}"); para(doc, "Status of the prior actions cannot be established from public information; it will be asked for in the inside-out stage.", size=8, italic=True, color=GREY)

    # ================= BACK ON TRACK =================
    bot = run.get("back_on_track") or []
    if bot:
        section(doc, "Getting back on track", "Actions behind or not started: what happened and how to recover")
        tb = table(doc, ["Action", "What happened", "Impact", "Root cause", "Mechanism to prevent", "Measures to track"], [3.0, 3.1, 2.6, 2.9, 3.0, 2.8], hsize=7)
        for b in bot:
            row = new_row(tb); set_cell(row[0], b.get("heading", b.get("id")), bold=True, color=NAVY, size=7.5)
            for k, key in enumerate(("what_happened", "impact", "root_cause", "prevention", "measures"), 1):
                v = b.get(key, ""); set_cell(row[k], "; ".join(v) if isinstance(v, list) else v, size=7.5)

    # ================= ADDITIONAL INFO =================
    ain = run.get("additional_information_needed") or []
    if ain:
        section(doc, "Additional information that would sharpen the assessment", "Only where it would change a grade or an action")
        for b in ain:
            p = doc.add_paragraph(); rich(p, "• " + b, 8.5); p.paragraph_format.space_after = Pt(1)

    # ================= COVERAGE AND LIMITATIONS =================
    section(doc, "Coverage and limitations", "What this assessment could and could not see")
    cvl = C["coverage"][lenses[-1]]; nac = cvl.get("not_assessable") or {}; cells = {i: (comp["indicators"].get(i) or {}).get(lenses[-1]) or {} for i in IND_IDS}
    confs = [c.get("confidence") for c in cells.values() if c.get("score") is not None]
    para(doc, f"{cvl['assessed']} of {cvl['of']} indicators assessed. Not assessable: " + " · ".join(f"{nac.get(k, 0)} {k}" for k in ("disclosure-limited", "retrieval-limited", "rule-limited")) + (f" · {nac['unspecified']} unspecified" if nac.get("unspecified") else "") + ".", size=8.5)
    para(doc, "Confidence of the assessed indicators: " + " · ".join(f"{sum(1 for c in confs if c == lv)} {lv}" for lv in ("High", "Medium", "Low")) + ".", size=8.5)
    nas = [f"{i} {IND_NAME[i]} ({c.get('na_cause') or 'cause not recorded'})" for i, c in cells.items() if c.get("score") is None]
    if nas: para(doc, "Not assessable, with cause: " + "; ".join(nas) + ".", size=8, after=3)
    cited = {x for c in comp["indicators"].values() for l in lenses for x in ((c.get(l) or {}).get("evidence") or [])}
    docs_ = {}  # the company's distinct cited documents, grouped by origin (peer sources are in Sources consulted); self-assessment items are excluded
    for k, e in run["evidence"].items():
        if k in cited and not is_sa(e) and e.get("document"): docs_.setdefault("Board documents (inside-out)" if e.get("evidence_class") == "inside_out" or e.get("source_origin") == "inside_out" else "Public record", set()).add(e["document"])
    kicker(doc, "Documents available")
    for g, ds in sorted(docs_.items(), reverse=True):
        para(doc, f"{g}: {len(ds)} document(s)", size=8, bold=True, color=NAVY, after=1, keep=True)
        for x in sorted(ds): para(doc, "• " + x, size=7.5, after=0)
    kicker(doc, "Documents unavailable"); un = (run.get("coverage") or {}).get("documents_unavailable") or []
    for x in un: para(doc, "• " + x, size=7.5, after=0)
    if not un: para(doc, "None recorded.", size=8, italic=True, color=GREY)
    if (run.get("coverage") or {}).get("notes"): kicker(doc, "Notes"); para(doc, run["coverage"]["notes"], size=8)

    # ================= GUARDRAILS =================
    section(doc, "Guardrails", "Guardrails")
    from app.domain.view import _guardrails  # one source with the dashboard and the React report
    groups, intro = _guardrails()
    if intro: para(doc, intro, size=8.5, after=4)
    for g in groups:
        kicker(doc, g["group"])
        for x in g["items"]: para(doc, "• " + x, size=8, after=2)
    doc.add_page_break()

    # ================= APPENDICES =================
    appn = 0
    def app_title(t): 
        nonlocal appn; appn += 1; section(doc, f"Appendix {appn}", t, 13)
    def sa_themes_appendix():
        # Anonymity: only synthesised themes are ever printed, never an answer or a per-respondent structure.
        app_title("Director self-assessment: themes")
        n = run.get("self_assessment_responses"); themes = run.get("self_assessment_themes") or []
        had = n or any(is_sa(e) for e in run["evidence"].values()) or any((run.get("self_assessment_answers") or {}).values())
        if not had: para(doc, "The inside-out view was prepared from board documents only, without a director self-assessment.", size=8.5); return
        para(doc, (f"{n} director response(s) were received. " if n else "Director responses were received; the number of respondents is not recorded in this run. ") + "Responses are synthesised into themes: no answer is quoted and nothing is attributed to an individual.", size=8, italic=True, color=GREY)
        if not themes: para(doc, "No synthesised themes were produced for this run.", size=8.5)
        for d in DIMS:
            ts = [t_ for t_ in themes if str(t_.get("dimension")) == d]
            if not ts: continue
            kicker(doc, f"{d}. {FW['dimensions'][d]}")
            for t_ in ts:
                p = doc.add_paragraph(); rich(p, t_.get("theme", ""), 8.5, bold=True, color=NAVY); p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
                para(doc, t_.get("summary", ""), size=8, after=3)
    app_title("Detailed indicator view")
    for d in DIMS:
        kicker(doc, f"{d}. {FW['dimensions'][d]}  ·  " + "  ·  ".join(f"{labels[l]} {C['dimensions'][d][l]['grade'] or 'n/a'}" for l in lenses))
        atext = (comp.get("assessment_text") or {}).get(d)
        if atext: para(doc, atext, size=8.5)
        hdr = ["Indicator"] + [labels[l] for l in lenses] + ["Rationale and sources", "Conf."]; widths = [3.2] + [1.0] * len(lenses) + [17.4 - 3.2 - 1.0 * len(lenses) - 1.6, 1.6]
        td = table(doc, hdr, widths)
        for i in [x for x in IND if x["dimension"] == d]:
            ind = comp["indicators"].get(i["id"], {}); row = new_row(td); set_cell(row[0], f"{i['id']} {i['name']}", bold=True, color=NAVY); k = 1
            for l in lenses: score_cell(row[k], (ind.get(l) or {}).get("score"), 9); k += 1
            rc = row[k]; rc.text = ""; first = True
            for l in lenses:
                cl = ind.get(l) or {}; p = rc.paragraphs[0] if first else rc.add_paragraph(); first = False
                lead = f"{labels[l]}: " if len(lenses) > 1 else None
                if cl.get("score") is None: rich(p, (lead or "") + (cl.get("rationale") or "Not assessable.") + (f" [{cl.get('na_cause')}]" if cl.get("na_cause") else ""), 7.5, italic=True, color=GREY)
                else: text_with_links(p, cl.get("rationale") or "", cl.get("evidence") or [], 7.5, lead)
            set_cell(row[k + 1], " / ".join((ind.get(l) or {}).get("confidence") or "–" for l in lenses), size=7.5, align="center")
        lim = (comp.get("limitations") or {}).get(d)
        if lim: para(doc, "Where this grade is limited: " + lim, size=7.5, italic=True, color=GREY)
    # indicator-level comparison: moved here from the body, which keeps to the narrative
    app_title("Indicator-level comparison")
    if io:
        kicker(doc, "Outside-in to post inside-out: indicators that moved")
        tc = table(doc, ["Indicator", "OI", "IO", "Δ", "Outside-in evidence, then what the board documents added"], [4.6, 1.0, 1.0, 0.9, 9.9])
        for i in IND_IDS:
            a = ((comp["indicators"].get(i) or {}).get("board") or {}).get("score"); b = ((comp["indicators"].get(i) or {}).get("board_io") or {}).get("score")
            if b is None or a == b: continue
            ind = next(x for x in IND if x["id"] == i); row = new_row(tc); set_cell(row[0], f"{i} {ind['name']}", bold=True, color=NAVY)
            score_cell(row[1], a, 8); score_cell(row[2], b, 8); set_cell(row[3], ("+" if a is None else ("+" if b > a else "")) + (str(b) if a is None else str(b - a)), align="center", bold=True, color=TEAL if a is None or b > a else "A6192E")
            p = row[4].paragraphs[0]; text_with_links(p, (run.get("io_change_notes") or {}).get(i) or first_sentence(((comp["indicators"][i].get("board_io") or {}).get("rationale"))), size=7.5)
    kicker(doc, "Evidence behind each indicator")
    hdr = ["Indicator"] + ([labels[l] for l in lenses] if io else ["Score"]) + ["Summary of the evidence"]; widths = [4.4] + [1.2] * len(lenses) + [17.4 - 4.4 - 1.2 * len(lenses)]
    te = table(doc, hdr, widths); summ = run.get("indicator_summary") or {}
    for i in IND:
        ind = comp["indicators"].get(i["id"], {}); row = new_row(te); set_cell(row[0], f"{i['id']} {i['name']}", bold=True, color=NAVY); k = 1
        for l in lenses: score_cell(row[k], (ind.get(l) or {}).get("score"), 8); k += 1
        p = row[k].paragraphs[0]; text_with_links(p, summ.get(i["id"]) or first_sentence((ind.get(lenses[-1]) or {}).get("rationale")) or "Not assessable.", size=7.5)
    # peer one-pagers
    appn += 1
    for n, p_ in enumerate(peers):
        pc = p_["computed"]
        section(doc, f"Appendix {appn}" + (f" ({n+1} of {len(peers)})"), f"{p_['name']}: outside-in indicator view", 13)
        if io: para(doc, "Carried forward from the outside-in report; peers are not re-assessed with inside-out information.", size=7.5, italic=True, color=GREY)
        gt = doc.add_table(rows=2, cols=7); borders(gt, "none"); fix_widths(gt, [2.48] * 7)
        for k, d in enumerate(DIMS): grade_cell(gt.rows[0].cells[k], pc["dimensions"][d]["board"]["grade"], 12); set_cell(gt.rows[1].cells[k], f"{d}. {FW['dimensions'][d]}", size=6.5, color=GREY, align="center")
        grade_cell(gt.rows[0].cells[6], pc["overall"]["board"].get("grade"), 12); set_cell(gt.rows[1].cells[6], "Overall", size=6.5, color=GREY, align="center", bold=True)
        if p_.get("summary"): para(doc, p_["summary"], size=8.5, before=4)
        tq = table(doc, ["Indicator", "Score", "Rationale and sources"], [4.4, 1.1, 11.9])
        for i in IND:
            cl = (p_["indicators"].get(i["id"]) or {}).get("board") or {}; row = new_row(tq); set_cell(row[0], f"{i['id']} {i['name']}", bold=True, color=NAVY, size=7.5); score_cell(row[1], cl.get("score"), 8)
            p = row[2].paragraphs[0]
            if cl.get("score") is None: rich(p, cl.get("rationale") or "Not assessable.", 7, italic=True, color=GREY)
            else: text_with_links(p, cl.get("rationale") or "", cl.get("evidence") or [], 7)
    if io: sa_themes_appendix()
    # how to read
    app_title("How to read this report")
    para(doc, "This assessment applies Jeff Bezos's Day 1 / Day 2 lens to a listed company's board, using " + ("public information only." if not io else "public information plus board documents and a director self-assessment supplied by the board.") + " Day 1 boards stay obsessed with the customer, make high-velocity decisions, resist proxies and process, embrace external trends, build rather than protect, and manage to cash. Day 2 is what creeps in afterwards.", size=8.5)
    para(doc, "Six dimensions, 22 indicators. Each indicator is scored 1 (Day 2), 3 (Mixed) or 5 (Day 1) against written anchors, once, for the board. Management's observable actions are attributed to the board's oversight, because a listed board is accountable for what the company does under its watch. Scores of 2 and 4 are used only where evidence meets parts of both adjacent anchors.", size=8.5)
    tg = table(doc, ["Average indicator score", "Grade", "Classification"], [5, 2, 10.4])
    for b in FW["grade_scale"]: row = new_row(tg); set_cell(row[0], FW["grade_scale_ranges"][b["grade"]]); grade_cell(row[1], b["grade"], 9); set_cell(row[2], b["label"])
    para(doc, "Overall grades are the rounded average of the dimension letters (A=5 … F=1; midpoints round down, so 2.5 is a D); provisional where only three dimensions are reported. Where evidence falls below the minimum, an indicator is marked not assessable and excluded rather than scored as Mixed, with the cause recorded (disclosure-limited, retrieval-limited, rule-limited). Peers are assessed on public information only.", size=8.5, before=4)
    app_title(f"Methodology and scoring rules (v{FW['rule_version']})")
    rules = ["Every indicator is scored once, for the board; a cell with nothing to assess is marked not applicable.", "Scores are 1, 3 or 5; 2 and 4 only where evidence meets parts of both adjacent anchors.", "A 1 or 5 requires two independent items from two documents, or one strong direct item (regulator or court record, formal board act, audited figure, or a board document); a 3 requires one.", "Confidence is assigned on the evidence: High = two direct sources or one strong source; Medium = one direct or two indirect; Low = a single indirect source, an inference, a silence-based score, or self-assessment alone.", "Below the minimum an indicator is not assessable and excluded, never defaulted to 3; silence is scorable in full on 1.4, 5.2, 5.3 and capped at 3 on 2.3, 2.4, 4.3; an empty register search supports up to 3 on 1.3.", "A dimension grade is reported only if at least half (outside-in) or two-thirds (post inside-out) of its indicators are assessed.", "Overall grade is the rounded average of dimension letters, with midpoints rounding down (for example, an average of 2.5 is a D); provisional at three dimensions; none below three.", "Behaviour indicators use 24 months (36 for 2.1 and 4.3); trend indicators five years of filings.", "3.2, 5.1, 6.1 and 6.2 are scored relative to the named peer set.", "Within scoped pairs an evidence item is primary for one indicator only.", "Approved news and trade press (for event facts) are direct evidence; formal board acts are strong direct evidence.", "Inside-out: board documents outrank public sources; the post inside-out score must cite at least one inside-out or self-assessment item; self-assessment alone is at most Medium confidence; peers are never re-scored.", "Each company is run twice on the same rule version; indicator scores within one point and unchanged dimension letters are required for the result to stand."]
    for k, rt in enumerate(rules, 1):
        p = doc.add_paragraph(); run_fmt(p.add_run(f"{k:02d}  "), 8, bold=True, color=BLUE); run_fmt(p.add_run(rt), 8); p.paragraph_format.space_after = Pt(2)
    # sources with bookmarks
    app_title("Sources consulted"); para(doc, "Each entry is referenced from the tables by its code. Verbatim excerpts from public and board documents are kept under 15 words; director self-assessment answers are never reproduced, because that would defeat their anonymity.", size=7.5, italic=True, color=GREY)
    used = {}
    for s in run["subjects"]:
        for iid, ind in (s.get("indicators") or {}).items():
            for l in ("board", "board_io"):
                for e in ((ind.get(l) or {}).get("evidence") or []): used.setdefault(e, set()).add(f"{iid} {IND_NAME.get(iid, '')}".strip() + ("" if s["role"] == "company" else f" ({s['name']})"))
    bid = 100
    for eid in sorted(run["evidence"], key=lambda x: (x[:2], len(x), x)):
        e = run["evidence"][eid]; p = doc.add_paragraph(); add_bookmark(p, f"EV_{eid}", bid); bid += 1
        run_fmt(p.add_run(f"{eid}  "), 8.5, bold=True, color=BLUE)
        if e.get("url"): add_external_link(p, e["url"], doc.part, 8.5, text=e.get("document") or e["url"], bold=True)
        else: run_fmt(p.add_run(e.get("document", "")), 8.5, bold=True, color=NAVY)
        run_fmt(p.add_run(f"   {e.get('date','')}   ·   {e.get('evidence_class','direct')}" + (f"   ·   {e.get('source_origin')}" if e.get("source_origin") else "") + ("   ·   SIMULATED" if e.get("simulated") else "") + ("   ·   RECENT" if e.get("recent") else "")), 7.5, color=GREY)
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
        if e.get("question"): para(doc, "Q: " + e["question"], size=7.5, italic=True, color=GREY, after=0)
        if is_sa(e): para(doc, "Responses are synthesised across directors; no individual answer is quoted (see the themes appendix).", size=7.5, color=GREY, after=0)
        elif e.get("excerpt"): para(doc, e["excerpt"], size=8, italic=True, after=0)
        if e.get("search_log"): para(doc, "Search: " + e["search_log"], size=7.5, color=GREY, after=0)
        p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        if e.get("url"): add_external_link(p2, e["url"], doc.part, 7)
        para(doc, "Used in: " + (", ".join(sorted(used.get(eid, []))) or "context"), size=7, color=GREY, after=5)
    # anchors
    app_title("Indicator definitions and anchors")
    for d in DIMS:
        kicker(doc, f"{d}. {FW['dimensions'][d]}")
        for i in [x for x in IND if x["dimension"] == d]:
            ov = (FW.get("outside_in_overrides") or {}).get(i["id"]); an = ov["anchors"] if (ov and not io) else i["anchors"]
            para(doc, f"{i['id']} {i['name']}", size=8.5, bold=True, color=NAVY, after=1, before=3, keep=True); para(doc, i["definition"] + ("  (Outside-in anchors applied.)" if ov and not io else ""), size=7.5, after=1)
            ta2 = table(doc, ["1 – Day 2", "3 – Mixed", "5 – Day 1"], [5.8, 5.8, 5.8], hsize=7); row = new_row(ta2)
            for k in range(3): set_cell(row[k], an[str([1, 3, 5][k])], size=7)
            para(doc, "Board prompt: " + i.get("board_lens", ""), size=7, color=GREY, after=3)
    doc.save(out_path)
    return out_path




# ---------------------------------------------------------------- HTML dashboard
def build_dashboard(run: dict, out_path: str) -> str:
    """Self-contained HTML dashboard (inline CSS/JS, no external assets); returns out_path."""
    run = _prep(run); comp = company_of(run)
    for e in (run.get("evidence") or {}).values():  # the page source must not carry a director's own words
        if isinstance(e, dict) and e.get("evidence_class") == "self_assessment": e["excerpt"] = ""
    run.pop("self_assessment_answers", None)
    tpl = (_HERE / "templates" / "dashboard.html").read_text(encoding="utf-8")
    page = (tpl.replace("__TITLE__", html.escape(f"{comp['name']} - Day 1 / Day 2 Board Assessment"))
            .replace("__RUN__", json.dumps(run).replace("</", "<\\/")).replace("__FW__", json.dumps(FW).replace("</", "<\\/")))
    Path(out_path).write_text(page, encoding="utf-8")
    return out_path


# ---------------------------------------------------------------- questionnaire


def _qshade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)


def _light_borders(table, color="BFC7D1"):
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color); b.append(e)
    table._tbl.tblPr.append(b)


def _qp(doc, text, size=10, bold=False, color=None, italic=False, after=6):
    para_ = doc.add_paragraph(); r = para_.add_run(clean(text)); r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"
    if color: r.font.color.rgb = rgb(color)
    para_.paragraph_format.space_after = Pt(after); return para_


SIM = {
 "1": ["Yes, a lost-contract debrief in Q2 changed our pricing posture; it reached us through the CEO's summary rather than directly. Most customer input arrives via management.", "Customer satisfaction sits in the annual bonus at about 10%; it has never reduced a payout to my knowledge."],
 "2": ["We revised our volume assumptions after a guidance miss last year; outside pressure prompted it, although the board had asked questions earlier.", "I trust the pipeline conversion number least. We have not retired a report in two years; the pack has grown."],
 "3": ["Electrification and channel shift. We hold a position on the first with capex targets; the second is discussed but without targets. Two external speakers in the past year.", "On tariffs we moved with peers; on the channel shift, after them. Management initiated both."],
 "4": ["Informally yes. A pricing pilot took one meeting; the divestiture took four. We deferred one acquisition for a further independent view.", "Minority views are not formally recorded. Decisions generally stay settled; one was revisited twice."],
 "5": ["Cost first, then buybacks to support EPS. Reinvestment was raised by two directors but not developed as an alternative.", "The last evaluation added a technology skills matrix; the last appointment brought digital experience. Owners would say we are slow to refresh."],
 "6": ["More time on earnings. No formal trigger; we would act on a rating warning. We refinanced early in 2025 before stress.", "Leverage policy is 2.0-2.5x net debt/EBITDA, reviewed annually, stress-tested at budget. Working capital is owned by the audit committee chair."]}



def build_questionnaire(company: str, out_path: str, simulate: bool = False) -> str:
    """Director self-assessment questionnaire (.docx) with editable answer boxes; simulate=True pre-fills synthetic answers."""
    out = out_path; sim = simulate; rnd = random.Random(7)
    company = clean(company or "the company")
    doc = Document(); st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2)
    _qp(doc, "ALVAREZ & MARSAL  |  DAY 1 / DAY 2 BOARD ASSESSMENT", 8, True, BLUE)
    _qp(doc, f"Director self-assessment: {company}", 18, True, NAVY, after=2)
    _qp(doc, "Twelve questions, two per dimension, about 30 minutes. Answer in the shaded boxes in your own words; specific examples matter more than complete sentences. Your answers are read alongside the outside-in evidence and treated as self-reported. Nothing you write is attributed to you by name in the output." + ("  THIS COPY CONTAINS SIMULATED ANSWERS FOR PROTOTYPE TESTING." if sim else ""), 9, color=GREY, after=10)
    meta = doc.add_table(rows=2, cols=2); _light_borders(meta)
    for r_, (k, v) in enumerate((("Your role on the board (optional)", ""), ("Date completed", ""))):
        meta.rows[r_].cells[0].text = k; meta.rows[r_].cells[1].text = v if not sim else ("Independent director (SIMULATED)" if r_ == 0 else "2026-10-23 (SIMULATED)")
        _qshade(meta.rows[r_].cells[0], TINT)
    doc.add_paragraph()
    n = 0; cur = None
    for d, dname, q in FW["questionnaire"]:
        if d != cur:
            cur = d; _qp(doc, f"Dimension {d}: {dname}", 12, True, NAVY, after=2)
        n += 1
        _qp(doc, f"Q{n}. {q}", 10, True, after=3)
        t = doc.add_table(rows=1, cols=1); _light_borders(t); cell = t.rows[0].cells[0]; _qshade(cell, "F7F9FB")
        cell.text = ""
        ans = SIM[d][(n - 1) % 2] if sim else ""
        rr = cell.paragraphs[0].add_run(f"[SIMULATED] {ans}" if sim else "Your answer:"); rr.font.size = Pt(9); rr.font.italic = not sim; rr.font.name = "Arial"
        if not sim: rr.font.color.rgb = rgb("9AA3B2")
        for _ in range(4 if not sim else 0): cell.add_paragraph()
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        if n % 2 == 0:
            rt = doc.add_table(rows=1, cols=len(RATING_LABELS) + 1); _light_borders(rt)
            rt.rows[0].cells[0].text = f"{dname} rating (optional): mark one"
            _qshade(rt.rows[0].cells[0], TINT)
            pick = rnd.choice([2, 3, 3, 4])  # one rating per dimension (legacy drew once per cell)
            for k, lab in enumerate(RATING_LABELS, 1):
                rt.rows[0].cells[k].text = ("☒ " if sim and k == pick else "☐ ") + lab
            doc.add_paragraph()
    _qp(doc, "Thank you. Return this file to the A&M team. " + DRAFT, 8, italic=True, color=GREY)
    doc.save(out)
    return out_path




def parse_questionnaire(path: str, run_date: str = "") -> dict:
    """Read a returned questionnaire (.docx). -> {answers: [str], ratings: {dim: 1-5}, simulated: bool,
    evidence: {SA01: {...}}} where evidence items are ready to merge into run["evidence"]."""
    doc = Document(path); qs = FW["questionnaire"]; answers: list[str] = []; ratings: dict = {}; qi = 0
    for t in doc.tables[1:]:  # skip the meta table
        if len(t.columns) == 1 and qi < len(qs):
            txt = "\n".join(p_.text for p_ in t.rows[0].cells[0].paragraphs).strip()
            if txt.startswith("Your answer:"): txt = ""
            answers.append(txt); qi += 1
        elif len(t.columns) == len(RATING_LABELS) + 1 and qi:
            d = qs[qi - 1][0]
            for k in range(1, len(RATING_LABELS) + 1):
                if "☒" in t.rows[0].cells[k].text: ratings[d] = k
    sim = any("[SIMULATED]" in a for a in answers); evidence = {}
    for n, ((d, dname, q_), a) in enumerate(zip(qs, answers), 1):
        if not a: continue
        evidence[f"SA{n:02d}"] = {"document": f"Director self-assessment, Q{n} (Dimension {d}: {dname})", "doc_type": "self-assessment", "date": run_date,
                                  "url": "", "excerpt": a[:200], "evidence_class": "self_assessment", "source_origin": "self_assessment", "simulated": sim, "question": q_}
    return {"answers": answers, "ratings": ratings, "simulated": sim, "evidence": evidence}



def file_names(short: str, oi: dict, io: dict | None, prior: dict | None) -> dict:
    """Deliverable file names, same keys/names as the legacy pipeline."""
    s = re.sub(r"[^A-Za-z0-9]+", "_", short or "").strip("_") or "Company"
    per = lambda r: (r or {}).get("run", {}).get("period", "")
    n = {"oi_docx": f"{s}_{per(oi)}_Outside-in_Report.docx", "oi_html": f"{s}_{per(oi)}_Outside-in_Dashboard.html",
         "quest": f"{s}_Self_Assessment_Questionnaire.docx"}
    if io: n.update(io_docx=f"{s}_{per(io)}_Inside-out_Update.docx", io_html=f"{s}_{per(io)}_Inside-out_Dashboard.html")
    if prior: n["prior_docx"] = f"{s}_{per(prior)}_Outside-in_Report.docx"
    return n

