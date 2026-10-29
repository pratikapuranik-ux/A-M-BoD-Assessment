"""Text extraction: docx, pdf, xlsx, text, limits and the PDF page budget."""
import base64
import io

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.services import documents as X


def make_docx():
    d = Document(); d.add_paragraph("Board charter"); t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Quorum"; t.rows[0].cells[1].text = "Five"
    b = io.BytesIO(); d.save(b); return b.getvalue()


def make_pdf(pages=2):
    b = io.BytesIO(); c = canvas.Canvas(b)
    for k in range(pages): c.drawString(72, 750, f"Minutes page {k + 1}"); c.showPage()
    c.save(); return b.getvalue()


def make_xlsx():
    wb = Workbook(); ws = wb.active; ws.title = "KPI"; ws.append(["metric", 4.5]); b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_docx_keeps_table_order():
    r = X.extract("c.docx", make_docx(), slot=3)
    assert r["text"] == "Board charter\nQuorum | Five" and r["slot"] == 3 and r["pdf_b64"] is None


def test_pdf_pages_and_b64():
    data = make_pdf(2); r = X.extract("m.pdf", data)
    assert r["pages"] == 2 and "[page 1]\nMinutes page 1" in r["text"] and "[page 2]" in r["text"]
    assert base64.b64decode(r["pdf_b64"]) == data and r["size"] == len(data)


def test_xlsx_and_text_and_mime():
    assert X.extract("k.xlsx", make_xlsx())["text"] == "[sheet KPI]\nmetric | 4.5"
    assert X.extract("n.txt", "café".encode("cp1252"))["text"] == "café"
    assert X.ext_of("noext", "application/pdf; charset=x") == ".pdf" and X.ext_of("A.DOCX") == ".docx"
    assert X.extract("noext", b"hello", mime="text/plain")["text"] == "hello"


def test_errors():
    with pytest.raises(X.ExtractError, match="unsupported"): X.extract("a.exe", b"x")
    with pytest.raises(X.ExtractError, match="no readable text"): X.extract("e.txt", b"   ")
    with pytest.raises(X.ExtractError, match="could not be read"): X.extract("bad.pdf", b"not a pdf")
    with pytest.raises(X.ExtractError, match="25 MB"): X.extract("big.txt", b"x" * (X.MAX_FILE + 1))
    with pytest.raises(X.ExtractError, match="empty"): X.extract_b64({"name": "z.txt", "data_b64": ""})


def test_extract_b64():
    r = X.extract_b64({"name": "d.docx", "slot": "2", "data_b64": base64.b64encode(make_docx()).decode()})
    assert r["slot"] == 2 and "Board charter" in r["text"]


def test_fit_pdfs_budget():
    small = {"name": "s.pdf", "size": 10, "pages": 60, "pdf_b64": "x", "text": "[page 1]\nhello"}
    big = {"name": "b.pdf", "size": 20, "pages": 60, "pdf_b64": "x", "text": "[page 1]\nbody text"}
    X.fit_pdfs([small, big])
    assert small["pdf_b64"] == "x" and big["pdf_b64"] is None  # over 100 pages: sent as text
    scan = {"name": "scan.pdf", "size": 30, "pages": 60, "pdf_b64": "x", "text": "[page 1]\n\n\n[page 2]\n"}
    with pytest.raises(X.ExtractError, match="scanned PDF"): X.fit_pdfs([dict(small), scan])
