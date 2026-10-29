"""Deliverable generators: build from the example runs, re-open and assert content."""
import json
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.services import deliverables as D
from app.services.pdf_summary import build_pdf_summary

EX = next(p for p in Path(__file__).resolve().parents if (p / "skill" / "example_runs").is_dir()) / "skill" / "example_runs"


def load(name):
    return json.loads((EX / f"{name}_run.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module", params=["outside_in", "inside_out"])
def run(request):
    return load(request.param)


def doc_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_docx(run, tmp_path):
    out = D.build_docx(run, str(tmp_path / "r.docx"))
    txt = doc_text(out)
    comp = D.company_of(run); io = run["run"]["stage"] == "inside_out"
    assert comp["name"] in txt
    for s in ("Executive summary", "Scorecard and trend", "Peer comparison", "Appendix 1", "Sources consulted", "Indicator definitions and anchors"):
        assert s in txt
    assert "DRAFT ONLY" in txt and "qualified Alvarez and Marsal professional" in txt
    assert comp["computed"]["overall"]["board"]["grade"] in txt
    for p in D.peers_of(run): assert p["name"] in txt
    assert ("Inside-out update" in txt) == io
    first = next(iter(run["evidence"]))
    assert first in txt
    assert Document(out).sections[0].footer.paragraphs[0].text.startswith("▲ DRAFT ONLY")


def test_removed_and_added_sections(run, tmp_path):
    txt = doc_text(D.build_docx(run, str(tmp_path / "s.docx"))); io = run["run"]["stage"] == "inside_out"
    for gone in ("Metrics worth tracking", "Data sources: public", "Simulated?", "Consistency run"): assert gone not in txt
    assert "Coverage and limitations" in txt and "Documents available" in txt and "Documents unavailable" in txt
    assert "Guardrails" in txt  # TODO: assert the authoritative guardrail text once it is supplied
    assert "Indicator-level comparison" in txt and "Evidence behind each indicator" in txt
    d = Document(str(tmp_path / "s.docx")); heads = [p.text for p in d.paragraphs]
    assert heads.index("Evidence behind each indicator") > heads.index("Appendix 1")
    assert ("Outside-in to post inside-out: indicators that moved" in heads) == io
    if io: assert heads.index("Outside-in to post inside-out: indicators that moved") > heads.index("Appendix 1")
    order = [heads.index(h) for h in ("Executive summary", "Peer comparison", "Coverage and limitations", "Guardrails", "Appendix 1")]
    assert order == sorted(order)


def test_no_verbatim_self_assessment(tmp_path):
    r = load("inside_out"); sa = {k: v for k, v in r["evidence"].items() if v.get("evidence_class") == "self_assessment"}
    assert sa
    for v in sa.values(): v["excerpt"] = "ZQXJ-secret-director-answer " + v["excerpt"]
    r["self_assessment_themes"] = [{"dimension": "1", "theme": "Customer signals reach the board late", "summary": "Directors describe customer input arriving mostly through management."}]
    txt = doc_text(D.build_docx(r, str(tmp_path / "sa.docx")))
    assert "ZQXJ-secret-director-answer" not in txt
    for v in sa.values(): assert v["excerpt"][:40] not in txt
    assert "Director self-assessment: themes" in txt and "Customer signals reach the board late" in txt
    assert "Responses are synthesised" in txt and "Verbatim excerpts from public and board documents" in txt
    r.pop("self_assessment_themes"); txt = doc_text(D.build_docx(r, str(tmp_path / "sa2.docx")))
    assert "Director self-assessment: themes" in txt and "No synthesised themes" in txt
    for k in list(r["evidence"]):
        if k in sa: del r["evidence"][k]
    r["self_assessment_answers"] = {}
    assert "board documents only, without a director self-assessment" in doc_text(D.build_docx(r, str(tmp_path / "sa3.docx")))


def test_questionnaire_names_every_dimension(tmp_path):
    import re
    txt = doc_text(D.build_questionnaire("Acme Corp", str(tmp_path / "q.docx")))
    assert not re.search(r"Dimension \d(?!:)", txt) and "Customer Obsession rating (optional)" in txt


def test_inside_out_body_shows_arrow_not_pair(tmp_path):
    d = Document(D.build_docx(load("inside_out"), str(tmp_path / "m.docx")))
    t = next(t for t in d.tables if t.rows[0].cells[0].text == "Dimension")
    assert not any(c.text.endswith(("Outside-in", "Inside-out")) or "OI→IO" in c.text for c in t.rows[0].cells)
    assert any(("▲" in c.text or "▼" in c.text) for r_ in t.rows[1:] for c in r_.cells)


def test_outside_in_without_prior_quarter_has_no_trend_furniture(tmp_path):
    """First-quarter outside-in report: no Trend row on the cover, 'SCORECARD' (not 'and trend'), no Change column, a Coverage line."""
    r = load("outside_in"); r["history"] = []
    out = tmp_path / "oi.docx"; D.build_docx(r, str(out)); d = Document(str(out))
    text = doc_text(str(out))
    assert "No prior report supplied" not in text and "Scorecard and trend" not in text and "What moved the rating" not in text
    assert "What drives the rating" in text and "Coverage: 22 of 22 indicators assessed." in text
    assert not any(t.rows[0].cells[0].text == "Dimension" and "Change" in [c.text for c in t.rows[0].cells] for t in d.tables)


def test_docx_strips_control_chars(tmp_path):
    r = load("outside_in"); r["headline"] = "Bad\x0bchar headline"
    txt = doc_text(D.build_docx(r, str(tmp_path / "c.docx")))
    assert "Badchar headline" in txt
    assert "\x0b" not in r["headline"].replace("\x0b", "") and r["headline"] == "Bad\x0bchar headline"  # input not mutated


def test_dashboard(run, tmp_path):
    out = D.build_dashboard(run, str(tmp_path / "d.html"))
    html = Path(out).read_text(encoding="utf-8")
    comp = D.company_of(run)
    assert html.startswith("<!DOCTYPE html>") and comp["name"] in html
    assert "__RUN__" not in html and "__FW__" not in html and "__TITLE__" not in html
    assert "DRAFT ONLY" in html
    assert "<script src" not in html and "<link " not in html and "https://cdn" not in html
    start = html.index("const RUN=") + len("const RUN="); end = html.index(", FW=", start)
    assert json.loads(html[start:end].replace("<\\/", "</"))["run"]["period"] == run["run"]["period"]


def test_questionnaire_roundtrip(tmp_path):
    blank = D.build_questionnaire("Acme Corp", str(tmp_path / "q.docx"))
    txt = doc_text(blank)
    assert "Director self-assessment: Acme Corp" in txt and "Q12." in txt and "DRAFT ONLY" in txt
    res = D.parse_questionnaire(blank)
    assert len(res["answers"]) == 12 and not any(res["answers"]) and res["evidence"] == {} and not res["simulated"]
    sim = D.build_questionnaire("Acme Corp", str(tmp_path / "qs.docx"), simulate=True)
    res = D.parse_questionnaire(sim, run_date="2026-10-23")
    assert res["simulated"] and len(res["evidence"]) == 12 and res["evidence"]["SA01"]["source_origin"] == "self_assessment"
    assert set(res["ratings"]) == {"1", "2", "3", "4", "5", "6"} and set(res["ratings"].values()) <= {1, 2, 3, 4, 5}


def test_file_names():
    oi, io, prior = load("outside_in"), load("inside_out"), {"run": {"period": "2026-Q3"}}
    n = D.file_names("Acme & Co.", oi, io, prior)
    p = oi["run"]["period"]
    assert n["oi_docx"] == f"Acme_Co_{p}_Outside-in_Report.docx" and n["oi_html"].endswith("Outside-in_Dashboard.html")
    assert n["quest"] == "Acme_Co_Self_Assessment_Questionnaire.docx" and n["prior_docx"].startswith("Acme_Co_2026-Q3")
    assert n["io_docx"].endswith("Inside-out_Update.docx") and "io_html" in n
    assert set(D.file_names("", oi, None, None)) == {"oi_docx", "oi_html", "quest"}


def test_pdf_summary(run, tmp_path):
    out = build_pdf_summary(run, str(tmp_path / "s.pdf"))
    r = PdfReader(out)
    assert 1 <= len(r.pages) <= 3
    txt = "\n".join(p.extract_text() for p in r.pages)
    comp = D.company_of(run)
    assert comp["name"] in txt and "SCORECARD" in txt and "ACTIONS" in txt and "DRAFT ONLY" in txt
    assert comp["computed"]["overall"]["board"]["label"].split()[0] in txt
    assert run["actions"][0]["heading"].split()[0] in txt


def test_ungraded_run_rejected(tmp_path):
    r = load("outside_in")
    del D.company_of(r)["computed"]
    try:
        import app.domain.engine  # noqa: F401
    except Exception:
        with pytest.raises(ValueError):
            D.build_docx(r, str(tmp_path / "x.docx"))


# ---------------------------------------------------------------- reconciliation, evidence links, appendix flow
REC = {"models": {"A": {"model": "claude-opus-5", "effort": "high"}, "B": {"model": "claude-sonnet-5-5", "effort": "high"}},
       "agreed": 80, "total": 88,
       "rows": [{"subject": "Chevron", "role": "company", "indicator": "1.4", "name": "Customer-linked accountability", "a": 3, "b": 4, "final": 3, "chosen": "A", "reason": "Run A cites [E01] directly."},
                {"subject": "BP", "role": "peer", "indicator": "2.1", "name": "Decision speed", "a": None, "b": 5, "final": 5, "chosen": "B", "reason": "See [E02] and [Z99]."}]}


def with_rec(name, rec=REC):
    r = load(name)
    if rec is not None: r["reconciliation"] = rec
    return r


@pytest.fixture(params=["outside_in", "inside_out"])
def stage_name(request):
    return request.param


def _recon_table(d):
    return next((t for t in d.tables if [c.text for c in t.rows[0].cells] == ["Cell", "Score (Run A / Run B)", "Final score", "Reason"]), None)


def test_reconciliation_table(stage_name, tmp_path):
    d = Document(D.build_docx(with_rec(stage_name), str(tmp_path / "r.docx")))
    t = _recon_table(d)
    assert t is not None and len(t.rows) == 3
    assert "Run A (claude-opus-5, high effort) and Run B (claude-sonnet-5-5, high effort) agreed on 80 of 88" in "\n".join(p.text for p in d.paragraphs)
    assert t.rows[1].cells[0].text.startswith("1.4 Customer-linked accountability") and t.rows[1].cells[1].text == "3 / 4"
    assert "Run A's score stands" in t.rows[1].cells[2].text
    assert "BP" in t.rows[2].cells[0].text and t.rows[2].cells[1].text == "n/a / 5"
    assert t.rows[0]._tr.trPr is not None and t.rows[0]._tr.trPr.find(D.qn("w:tblHeader")) is not None
    heads = [p.text for p in d.paragraphs]
    assert heads.index("Where the two runs differed") > max(i for i, h in enumerate(heads) if h.startswith("Dimension grades against the peer set"))
    assert heads.index("Where the two runs differed") < heads.index(next(h for h in heads if h.startswith("What the board could do next")))


def test_reconciliation_empty_and_absent(stage_name, tmp_path):
    d = Document(D.build_docx(with_rec(stage_name, {**REC, "rows": [], "agreed": 88}), str(tmp_path / "e.docx")))
    assert "Both runs agreed on every indicator score; nothing needed reconciling." in "\n".join(p.text for p in d.paragraphs) and _recon_table(d) is None
    d = Document(D.build_docx(with_rec(stage_name, None), str(tmp_path / "a.docx")))
    assert "Where the two runs differed" not in "\n".join(p.text for p in d.paragraphs) and _recon_table(d) is None


def test_every_evidence_code_is_a_link(stage_name, tmp_path):
    run = with_rec(stage_name); out = D.build_docx(run, str(tmp_path / "l.docx")); d = Document(out)
    known = set(run["evidence"])
    body = d.element.body; W = D.qn
    codes = linked = 0
    for p in body.iter(W("w:p")):
        in_link = [h for h in p.iter(W("w:hyperlink"))]
        link_text = "".join(t.text or "" for h in in_link for t in h.iter(W("w:t")))
        full = "".join(t.text or "" for t in p.iter(W("w:t")))
        for m in D.EID_RE.finditer(full):
            if m.group(1) in known: codes += 1
        linked += sum(1 for m in D.EID_RE.finditer(link_text) if m.group(1) in known)
        for h in in_link:
            if h.get(W("w:anchor")): assert h.get(W("w:anchor")).startswith("EV_")
    # the bookmarks themselves sit in the sources appendix; each code in the text outside a bookmark entry must be linked
    assert codes > 20 and linked == codes
    assert "[Z99]" in "\n".join(c.text for c in _recon_table(d).rows[2].cells)  # no source entry: stays plain text
    bookmarks = {b.get(W("w:name")) for b in body.iter(W("w:bookmarkStart"))}
    assert all(h.get(W("w:anchor")) in bookmarks for h in body.iter(W("w:hyperlink")) if h.get(W("w:anchor")))


def test_source_titles_link_to_urls(tmp_path):
    run = with_rec("outside_in"); d = Document(D.build_docx(run, str(tmp_path / "u.docx")))
    ext = {r.target_ref for r in d.part.rels.values() if r.is_external}
    urls = {e["url"] for e in run["evidence"].values() if e.get("url")}
    assert urls and urls <= ext


def test_no_page_breaks_inside_appendices(stage_name, tmp_path):
    d = Document(D.build_docx(with_rec(stage_name), str(tmp_path / "b.docx")))
    paras = d.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.upper() == "APPENDIX 1" or p.text == "Appendix 1")
    W = D.qn
    assert not any(br.get(W("w:type")) == "page" for p in paras[start:] for br in p._p.iter(W("w:br")))
    assert not any(p._p.find(".//" + W("w:pageBreakBefore")) is not None for p in paras[start:])
    before = sum(1 for p in paras[:start] for br in p._p.iter(W("w:br")) if br.get(W("w:type")) == "page")
    assert before == 3  # after the cover, after the executive summary, before Appendix 1
    assert all(p.paragraph_format.keep_with_next for p in paras[start:start + 2])
