"""Word helpers shared by build_docx.py and build_questionnaire.py (A&M style)."""
import re
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from common import *

EID_RE = re.compile(r"\[((?:[EXBKLN]|SA|IO)\d{2,3})\]")  # the company (E), peers (X B K L N), self-assessment, board documents
CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def clean(text): return CTRL_RE.sub("", text) if isinstance(text, str) else text

# python-docx rejects control characters (a pasted vertical tab in an answer would stop the report); every run goes through here
_add_run = Paragraph.add_run
Paragraph.add_run = lambda self, text=None, style=None: _add_run(self, clean(text), style)

def rgb(h): return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)

def borders(table, kind="light", color="D9DEE5"):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        if kind == "none": e.set(qn("w:val"), "nil")
        else: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)

def fix_widths(t, widths_cm):
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm): gc.set(qn("w:w"), str(int(w * 567)))
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None: tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    for row in t.rows:
        for i, w in enumerate(widths_cm): row.cells[i].width = Cm(w)
    t._widths = widths_cm

def new_row(t):
    r = t.add_row()
    for i, w in enumerate(getattr(t, "_widths", [])): r.cells[i].width = Cm(w)
    for c in r.cells: cell_margins(c)
    return r.cells

def run_fmt(r, size=9, bold=False, italic=False, color=None, caps=False, spacing=None):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"
    if color: r.font.color.rgb = rgb(color)
    if caps: r.font.all_caps = True
    if spacing is not None:
        rPr = r._r.get_or_add_rPr(); sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), str(spacing)); rPr.append(sp)

def para(host, text="", size=9.5, bold=False, italic=False, color=None, align=None, after=4, before=0, caps=False, spacing=None, keep=False):
    p = host.add_paragraph()
    if text: run_fmt(p.add_run(text), size, bold, italic, color, caps, spacing)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if keep: p.paragraph_format.keep_with_next = True
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def set_cell(cell, text, size=8, bold=False, color=None, align=None, italic=False):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    run_fmt(p.add_run("" if text is None else clean(str(text))), size, bold, italic, color)
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def kicker(doc, text): return para(doc, text, size=7.5, bold=True, color=BLUE, caps=True, spacing=40, after=1, before=8, keep=True)
def headline(doc, text, size=13): return para(doc, text, size=size, bold=True, color=NAVY, after=4, keep=True)
def section(doc, kick, head, size=13): kicker(doc, kick); headline(doc, head, size)

def grade_cell(cell, grade, size=10):
    set_cell(cell, grade or "n/a", size=size, bold=True, color="FFFFFF", align="center"); shade(cell, GRADE_FILL.get(grade))

def score_cell(cell, score, size=9):
    if score is None: set_cell(cell, "n/a", size=7, color=GREY, align="center"); shade(cell, LIGHT)
    else: set_cell(cell, score, size=size, bold=True, color="FFFFFF", align="center"); shade(cell, SCORE_FILL[score])

def table(doc, headers, widths, hsize=7.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, size=hsize, bold=True, color=NAVY); shade(t.rows[0].cells[i], "DCEBF7"); cell_margins(t.rows[0].cells[i])
    fix_widths(t, widths); return t

def add_bookmark(paragraph, name, bid):
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid)); paragraph._p.insert(0, s); paragraph._p.append(e)

def _link_run(text, size):
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), BLUE); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial"); rPr.append(rf)
    r.append(rPr); t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = clean(text); r.append(t); return r

def add_internal_link(paragraph, text, anchor, size=7.5):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1"); h.append(_link_run(text, size)); paragraph._p.append(h)

def add_external_link(paragraph, url, part, size=7):
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid); h.append(_link_run(url, size)); paragraph._p.append(h)

def text_with_links(paragraph, text, evidence_ids=None, size=7.5, lead=None, lead_color=NAVY):
    if lead: run_fmt(paragraph.add_run(lead), size, bold=True, color=lead_color)
    pos = 0; seen = set(); text = text or ""
    for m in EID_RE.finditer(text):
        if m.start() > pos: run_fmt(paragraph.add_run(text[pos:m.start()]), size)
        add_internal_link(paragraph, f"[{m.group(1)}]", f"EV_{m.group(1)}", size); seen.add(m.group(1)); pos = m.end()
    if text[pos:]: run_fmt(paragraph.add_run(text[pos:]), size)
    rest = [e for e in (evidence_ids or []) if e not in seen]
    if rest:
        run_fmt(paragraph.add_run(" Sources: "), size, italic=True, color=GREY)
        for k, e in enumerate(rest):
            if k: run_fmt(paragraph.add_run(", "), size)
            add_internal_link(paragraph, f"[{e}]", f"EV_{e}", size)
    paragraph.paragraph_format.space_after = Pt(1)

def add_page_number(paragraph):
    r = paragraph.add_run(); run_fmt(r, 8, color=GREY)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag: f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tag); r._r.append(f)
        else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)

def first_sentence(text):
    text = (text or "").strip(); f, _, _ = text.partition(". "); return (f.rstrip(".") + ".") if f else ""
