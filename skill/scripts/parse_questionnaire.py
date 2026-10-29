#!/usr/bin/env python3
"""
parse_questionnaire.py - read a returned self-assessment questionnaire (.docx) and append its answers
to the run JSON as self_assessment evidence items (SA01..SA12), plus any self-ratings.

Usage:
    python parse_questionnaire.py run.json filled_questionnaire.docx

The answers become evidence items with evidence_class "self_assessment", source_origin
"self_assessment", document "Director self-assessment, Q<n> (Dimension <d>)", and the first 200
characters of the answer as excerpt. Score indicators against them in the board_io cells by hand.
"""
import json, sys
from docx import Document
from common import *


def main():
    if len(sys.argv) < 3: print(__doc__); sys.exit(2)
    run = json.load(open(sys.argv[1])); doc = Document(sys.argv[2])
    qs = FW["questionnaire"]; answers = []; ratings = {}
    tables = doc.tables[1:]  # skip meta table
    qi = 0
    for t in tables:
        if len(t.columns) == 1 and qi < len(qs):
            txt = "\n".join(p_.text for p_ in t.rows[0].cells[0].paragraphs).strip()
            if txt.startswith("Your answer:"): txt = ""
            answers.append(txt); qi += 1
        elif len(t.columns) == 4:
            d = qs[qi - 1][0]
            for k, lab in ((1, 1), (2, 3), (3, 5)):
                if "☒" in t.rows[0].cells[k].text: ratings[d] = lab
    sim = any("[SIMULATED]" in a for a in answers)
    run.setdefault("evidence", {})
    for n, ((d, dname, q), a) in enumerate(zip(qs, answers), 1):
        if not a: continue
        run["evidence"][f"SA{n:02d}"] = {"document": f"Director self-assessment, Q{n} (Dimension {d}: {dname})", "doc_type": "self-assessment", "date": run["run"].get("run_date", ""),
            "url": "", "excerpt": a[:200], "evidence_class": "self_assessment", "source_origin": "self_assessment", "simulated": sim, "question": q}
    run["self_ratings"] = ratings
    json.dump(run, open(sys.argv[1], "w"), indent=1)
    print(f"{sum(1 for a in answers if a)} answers added as SA items; self-ratings: {ratings}" + (" [SIMULATED]" if sim else ""))


if __name__ == "__main__":
    main()
