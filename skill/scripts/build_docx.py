#!/usr/bin/env python3
"""
build_docx.py (v1.2, board-only) - render the Word report from a graded run JSON.

Usage:
    python build_docx.py run.json output.docx

Stage-aware. Outside-in report:
  Cover | page break | Executive summary | page break | Scorecard and trend · Evidence by indicator ·
  Peer comparison · Actions · Getting back on track · Additional information needed (optional) |
  page break | Appendix 1 Detailed indicator view · Appendix 2 Peer one-pagers · then How to read,
  Methodology, Sources consulted (bookmarked), Indicator anchors.
Inside-out report: same, with outside-in and post-inside-out grades side by side, inside-vs-outside
summary in the executive summary, peers carried forward with a note, and Appendix 3 = sources overview
(public / inside-out / simulated) after the peer one-pagers.
Main body is written to fit five pages (six with trend tables); only the three page breaks above are
used in the body; peer one-pagers in the appendix each start on a new page.
"""
import json, sys
from docx import Document
from docx.enum.text import WD_BREAK
from docx_helpers import *


def hist_grade(h, d=None):
    """A history entry's dimension (or overall) grade: board_io for an inside-out quarter, as build_dashboard does."""
    x = h["dimensions"][d] if d else h["overall"]
    return (x.get("board_io") or x.get("board") or {}).get("grade") if h["stage"] == "inside_out" else (x.get("board") or {}).get("grade")


def main():
    if len(sys.argv) < 3: print(__doc__); sys.exit(2)
    run = json.load(open(sys.argv[1])); out = sys.argv[2]
    comp = company_of(run); peers = peers_of(run); C = comp["computed"]
    stage = stage_of(run); lenses, labels = lenses_for(run); io = stage == "inside_out"
    period = period_label(run); hist = run.get("history", []); mode_label = FW["run_modes"].get(run["run"].get("mode"), "")
    simulated_note = " SIMULATED" if run["run"].get("simulated") else ""
    title = "Day 1 / Day 2 Board Assessment" + (": inside-out update" if io else "")

    doc = Document(); st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9.5); st.font.color.rgb = rgb("1A1F2B")
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.6)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; hp.text = ""; run_fmt(hp.add_run("ALVAREZ & MARSAL"), 7, bold=True, color=NAVY, spacing=30)
    run_fmt(hp.add_run(f"   |   {title}   |   {comp['name']}   |   {period}   |   {mode_label}{simulated_note}"), 7, color=GREY)
    fp = sec.footer.paragraphs[0]; fp.text = ""; run_fmt(fp.add_run("▲ " + DRAFT + "      "), 7, italic=True, color=AMBER); add_page_number(fp)

    # ================= COVER =================
    t = doc.add_table(rows=1, cols=1); borders(t, "none"); fix_widths(t, [17.4]); c = t.rows[0].cells[0]; shade(c, "F1F7FC"); cell_margins(c, 700, 700, 600, 600); c.text = ""
    r0 = c.paragraphs[0]; run_fmt(r0.add_run("ALVAREZ & MARSAL"), 9, bold=True, color=BLUE, spacing=60); r0.paragraph_format.space_after = Pt(60)
    para(c, "Day 1 / Day 2", size=30, bold=True, color=NAVY, after=0); para(c, "Board Assessment", size=30, bold=True, color=NAVY, after=6)
    if io: para(c, "Inside-out update", size=16, bold=True, color=AMBER, after=10)
    para(c, f"{comp['name']}" + (f" ({comp.get('ticker')})" if comp.get("ticker") else "") + f"  ·  {period}", size=12, color=NAVY, after=4)
    para(c, ("Outside-in assessment from public information" if not io else "Outside-in assessment updated with board documents and the director self-assessment") + f", benchmarked against {', '.join(p['name'] for p in peers)}", size=10, color="646464", after=30)
    bar = c.add_paragraph(); run_fmt(bar.add_run("────────────"), 10, color=BLUE); para(c, run["run"].get("run_date", ""), size=10, bold=True, color=NAVY, after=0)
    doc.add_paragraph()
    info = doc.add_table(rows=0, cols=2); borders(info, "none"); fix_widths(info, [3.2, 14.2])
    for k, v in (("Purpose", f"Grade the board of {comp['name']} A-F on six Day 1 / Day 2 dimensions and 22 indicators, attributing what the company did under the board's watch to the board, against a named peer set."),
                 ("Basis", ("Public information only." if not io else "Public information plus the board documents and self-assessment supplied by the director; peers remain on public information.") + f" Rules v{FW['rule_version']}. {mode_label}." + (" SIMULATED RUN." if run["run"].get("simulated") else "")),
                 ("Trend", (f"Compared with {len(hist)} prior quarter(s): {', '.join(h['period'] + (' (simulated)' if h['simulated'] else '') for h in hist)}." if hist else "No prior report supplied; no trend shown.")),
                 ("Evidence window", f"Behaviour indicators from {run['run'].get('document_window', {}).get('behaviour_from', '')}; trend indicators from {run['run'].get('document_window', {}).get('trend_from', '')}.")):
        row = new_row(info); set_cell(row[0], k, size=8, bold=True, color=NAVY); shade(row[0], TINT); set_cell(row[1], v, size=8.5)
    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1); borders(box, "none"); fix_widths(box, [17.4]); bc = box.rows[0].cells[0]; shade(bc, "FFF5E6"); cell_margins(bc, 120, 120, 200, 200); bc.text = ""
    run_fmt(bc.paragraphs[0].add_run("DRAFT ONLY — " + mode_label.upper()), 8, bold=True, color=AMBER)
    para(bc, "Grades apply to the board as a body; no individual is assessed. Where a grade rests on outside-in information, additional documentation can arrive at a more accurate grade. Requires review and approval by a qualified Alvarez and Marsal professional before use.", size=8, color="4A4A4A", after=0)
    doc.add_page_break()

    # ================= EXECUTIVE SUMMARY =================
    section(doc, "Executive summary", run.get("headline") or f"{comp['name']}: the board's Day 1 / Day 2 grade", 14)
    tiles = doc.add_table(rows=2, cols=len(lenses)); borders(tiles, "none"); fix_widths(tiles, [17.4 / len(lenses)] * len(lenses))
    for k, l in enumerate(lenses):
        o = C["overall"][l]; cell = tiles.rows[0].cells[k]; shade(cell, TINT); cell_margins(cell, 140, 140, 200, 200); cell.text = ""
        run_fmt(cell.paragraphs[0].add_run(labels[l].upper()), 7.5, bold=True, color=GREY, spacing=30)
        p = cell.add_paragraph(); run_fmt(p.add_run(o.get("grade") or "—"), 30, bold=True, color=GRADE_FILL.get(o.get("grade"))); p.paragraph_format.space_after = Pt(0)
        run_fmt(p.add_run("   " + o.get("label", "")), 9.5, bold=True, color=NAVY)
        if o.get("provisional"): run_fmt(p.add_run("   PROVISIONAL"), 7.5, bold=True, color=AMBER)
        prev = next((h for h in reversed(hist) if l in h.get("overall", {}) or l == "board"), None)
        pv = prev["overall"].get(l if l in prev["overall"] else "board", {}).get("grade") if prev else None
        dl = grade_delta(o.get("grade"), pv)
        para(cell, f"{o['dimensions_reported']} of {len(DIMS)} dimensions reported" + (f"  ·  prior quarter {pv} ({'+' if dl > 0 else ''}{dl} step{'s' if abs(dl) != 1 else ''})" if pv and dl is not None else ""), size=7.5, color=GREY, after=0)
        c2 = tiles.rows[1].cells[k]; c2.text = ""; cell_margins(c2, 60, 60, 200, 200); run_fmt(c2.paragraphs[0].add_run((run.get("overall_summary") or {}).get(l, "")), 8.5)
    kicker(doc, "What moved the rating"); para(doc, run.get("what_moved") or " ".join(run.get("headline_findings", [])), size=9)
    if io and run.get("inside_vs_outside"):
        ivo = run["inside_vs_outside"]; kicker(doc, "Inside view against the public record")
        t2 = table(doc, ["Better than the public evidence suggested", "Worse than the public evidence suggested"], [8.7, 8.7])
        row = new_row(t2)
        for k, key in enumerate(("better", "worse")):
            row[k].text = ""; first = True
            for item in ivo.get(key, []):
                p = row[k].paragraphs[0] if first else row[k].add_paragraph(); first = False; text_with_links(p, "• " + item, size=7.5)
    kicker(doc, "Key actions"); 
    for a in run.get("actions", [])[:5]:
        p = doc.add_paragraph(); run_fmt(p.add_run("▸ " + a["heading"] + ". "), 8.5, bold=True, color=NAVY); run_fmt(p.add_run(first_sentence(a["description"])), 8.5); p.paragraph_format.space_after = Pt(1)
    if run.get("key_metrics"):
        kicker(doc, "Metrics worth tracking"); t3 = table(doc, ["Metric", "Why it matters", "Now"], [5.2, 9.2, 3.0])
        for m in run["key_metrics"][:6]:
            row = new_row(t3); set_cell(row[0], m.get("metric"), bold=True, color=NAVY); set_cell(row[1], m.get("why")); set_cell(row[2], m.get("current", "") or "")
    kicker(doc, "Position relative to peers"); para(doc, run.get("peer_position", ""), size=9)
    for ln in run.get("peer_summary_lines", [])[:5]:
        p = doc.add_paragraph(); run_fmt(p.add_run("• " + ln), 8.5); p.paragraph_format.space_after = Pt(1)
    doc.add_page_break()

    # ================= SCORECARD AND TREND =================
    section(doc, "Scorecard and trend", "Grades by dimension" + (f", {len(hist)} prior quarter(s)" if hist else ""))
    hcols = [h["period"] for h in hist]
    hdr = ["Dimension"] + [f"{h['period']}{' (sim)' if h['simulated'] else ''}" for h in hist] + [f"{period} {labels[l]}" if io else period for l in lenses] + ["Change"] + ["Evidence in this run"]
    wfix = 3.6 + 1.1 * len(hist) + 1.3 * len(lenses) + 1.2; widths = [3.6] + [1.1] * len(hist) + [1.3] * len(lenses) + [1.2] + [17.4 - wfix]
    tt = table(doc, hdr, widths)
    for d in DIMS:
        row = new_row(tt); set_cell(row[0], f"{d}. {FW['dimensions'][d]}", bold=True, color=NAVY); k = 1
        for h in hist:
            grade_cell(row[k], hist_grade(h, d), 8.5); k += 1
        for l in lenses: grade_cell(row[k], C["dimensions"][d][l]["grade"], 9.5); k += 1
        cur = C["dimensions"][d][lenses[-1]]["grade"]; prevg = hist_grade(hist[-1], d) if hist else None  # an inside-out prior counts at its board_io grade
        dl = grade_delta(cur, prevg)
        set_cell(row[k], ("▲" if dl > 0 else "▼" if dl < 0 else "=") + (f" {abs(dl)}" if dl else "") if dl is not None else ("OI→IO " + ("▲" if grade_delta(C["dimensions"][d]["board_io"]["grade"], C["dimensions"][d]["board"]["grade"]) > 0 else "▼" if grade_delta(C["dimensions"][d]["board_io"]["grade"], C["dimensions"][d]["board"]["grade"]) < 0 else "=") if io and C["dimensions"][d]["board_io"]["grade"] and C["dimensions"][d]["board"]["grade"] else "–"), size=8, align="center", color=(TEAL if (dl or 0) > 0 else "A6192E" if (dl or 0) < 0 else GREY)); k += 1
        set_cell(row[k], (run.get("scorecard_headline") or {}).get(d, ""), size=7.5)
    row = new_row(tt); set_cell(row[0], "Overall", bold=True, color=NAVY); shade(row[0], "DCEBF7"); k = 1
    for h in hist: grade_cell(row[k], hist_grade(h), 8.5); k += 1
    for l in lenses: grade_cell(row[k], C["overall"][l].get("grade"), 9.5); k += 1
    set_cell(row[k], "", size=8); k += 1
    set_cell(row[k], "Rounded average of dimension letters (midpoints round down)." + (" * provisional" if any(C["overall"][l].get("provisional") for l in lenses) else ""), size=7, italic=True)
    cv = C["coverage"][lenses[-1]]; na_cells = [f"{i} ({((comp['indicators'].get(i) or {}).get(lenses[-1]) or {}).get('na_cause','')})" for i in IND_IDS if ((comp["indicators"].get(i) or {}).get(lenses[-1]) or {}).get("score") is None]
    para(doc, ("Grades rest on outside-in information; where an indicator is not assessable or scored on thin public evidence, board documents can arrive at a more accurate grade." if not io else "Post inside-out grades use the board documents and self-assessment supplied; outside-in grades are shown unchanged for comparison."), size=7.5, italic=True, color=GREY, before=2)
    # what changed table
    changes = []
    if hist:
        prev = hist[-1]
        for i in IND_IDS:
            cur = ((comp["indicators"].get(i) or {}).get("board") or {}).get("score"); pv = (prev["indicators"].get(i) or {}).get("board")
            if cur is not None and pv is not None and cur != pv: changes.append((i, pv, cur))
    if changes:
        kicker(doc, f"What changed since {hist[-1]['period']}" + (" (simulated prior)" if hist[-1]["simulated"] else ""))
        tc = table(doc, ["Indicator", "Prior", "Now", "Δ", "Evidence that moved the score"], [4.6, 1.1, 1.1, 0.9, 9.7])
        notes = run.get("change_notes") or {}
        for i, pv, cur in changes:
            ind = next(x for x in IND if x["id"] == i); row = new_row(tc); set_cell(row[0], f"{i} {ind['name']}", bold=True, color=NAVY)
            score_cell(row[1], pv, 8); score_cell(row[2], cur, 8); set_cell(row[3], f"{'+' if cur > pv else ''}{cur - pv}", align="center", bold=True, color=TEAL if cur > pv else "A6192E")
            p = row[4].paragraphs[0]; text_with_links(p, notes.get(i) or first_sentence(((comp["indicators"][i].get("board") or {}).get("rationale"))), size=7.5)
    if io:
        kicker(doc, "Outside-in to post inside-out: indicators that moved")
        tc = table(doc, ["Indicator", "OI", "IO", "Δ", "Outside-in evidence, then what the board documents added"], [4.6, 1.0, 1.0, 0.9, 9.9])
        for i in IND_IDS:
            a = ((comp["indicators"].get(i) or {}).get("board") or {}).get("score"); b = ((comp["indicators"].get(i) or {}).get("board_io") or {}).get("score")
            if b is None or a == b: continue
            ind = next(x for x in IND if x["id"] == i); row = new_row(tc); set_cell(row[0], f"{i} {ind['name']}", bold=True, color=NAVY)
            score_cell(row[1], a, 8); score_cell(row[2], b, 8); set_cell(row[3], ("+" if a is None else ("+" if b > a else "")) + (str(b) if a is None else str(b - a)), align="center", bold=True, color=TEAL if a is None or b > a else "A6192E")
            p = row[4].paragraphs[0]; text_with_links(p, (run.get("io_change_notes") or {}).get(i) or first_sentence(((comp["indicators"][i].get("board_io") or {}).get("rationale"))), size=7.5)

    # ================= EVIDENCE BY INDICATOR =================
    section(doc, "Evidence behind each indicator", "One line per indicator; the detailed view is in the appendix")
    hdr = ["Indicator"] + ([labels[l] for l in lenses] if io else ["Score"]) + ["Summary of the evidence"]; widths = [4.4] + [1.2] * len(lenses) + [17.4 - 4.4 - 1.2 * len(lenses)]
    te = table(doc, hdr, widths); summ = run.get("indicator_summary") or {}
    for i in IND:
        ind = comp["indicators"].get(i["id"], {}); row = new_row(te); set_cell(row[0], f"{i['id']} {i['name']}", bold=True, color=NAVY); k = 1
        for l in lenses: score_cell(row[k], (ind.get(l) or {}).get("score"), 8); k += 1
        p = row[k].paragraphs[0]; text_with_links(p, summ.get(i["id"]) or first_sentence((ind.get(lenses[-1]) or {}).get("rationale")) or "Not assessable.", size=7.5)

    # ================= PEER COMPARISON =================
    section(doc, "Peer comparison", "Dimension grades against the peer set" + (" (outside-in grades only)" if io else ""))
    if io: para(doc, "Peers are compared on outside-in grades only; their inside-out information is not available. The company's outside-in grade is used in this table.", size=7.5, italic=True, color=GREY)
    hdr = ["Dimension", comp["name"]] + [p_["name"] for p_ in peers] + ["How the company compares"]
    pw = min(1.5, 7.0 / (1 + len(peers))); widths = [2.9] + [pw] * (1 + len(peers)) + [17.4 - 2.9 - pw * (1 + len(peers))]
    tp = table(doc, hdr, widths, hsize=7); rs = run.get("peer_row_summary") or {}
    for d in DIMS:
        row = new_row(tp); set_cell(row[0], f"{d}. {FW['dimensions'][d]}", bold=True, color=NAVY, size=7.5); grade_cell(row[1], C["dimensions"][d]["board"]["grade"], 9)
        for k, p_ in enumerate(peers, 2): grade_cell(row[k], p_["computed"]["dimensions"][d]["board"]["grade"], 9)
        set_cell(row[len(peers) + 2], rs.get(d, (run.get("comparative_evidence") or {}).get(d, "")), size=7.5)
    row = new_row(tp); set_cell(row[0], "Overall", bold=True, color=NAVY, size=7.5); shade(row[0], "DCEBF7"); grade_cell(row[1], C["overall"]["board"].get("grade"), 9)
    for k, p_ in enumerate(peers, 2): grade_cell(row[k], p_["computed"]["overall"]["board"].get("grade"), 9)
    set_cell(row[len(peers) + 2], run.get("peer_position", ""), size=7.5)
    if hist and hist[-1].get("peers"):
        ph = hist[-1]["peers"]; kicker(doc, f"Peer trend: {hist[-1]['period']} to {period} (overall)")
        tr = table(doc, ["Company"] + [p_["name"] for p_ in peers], [4.0] + [(13.4) / len(peers)] * len(peers), hsize=7)
        row = new_row(tr); set_cell(row[0], hist[-1]["period"], bold=True)
        for k, p_ in enumerate(peers, 1): grade_cell(row[k], (ph.get(p_["name"]) or {}).get("overall"), 8.5)
        row = new_row(tr); set_cell(row[0], period, bold=True)
        for k, p_ in enumerate(peers, 1): grade_cell(row[k], p_["computed"]["overall"]["board"].get("grade"), 8.5)

    # ================= ACTIONS =================
    section(doc, "Actions", "What the board could do next" + (" (updated after the inside-out review)" if io else ""))
    for k, a in enumerate(run.get("actions", [])[:5], 1):
        p = doc.add_paragraph(); run_fmt(p.add_run(f"{k:02d}  "), 10, bold=True, color=BLUE); run_fmt(p.add_run(a["heading"]), 9.5, bold=True, color=NAVY); p.paragraph_format.space_after = Pt(1); p.paragraph_format.keep_with_next = True
        p = doc.add_paragraph(); text_with_links(p, a["description"], size=8.5); p.paragraph_format.space_after = Pt(1)
        if a.get("measures") or a.get("target_by"):
            para(doc, ("Measures: " + "; ".join(a.get("measures", [])) if a.get("measures") else "") + (f"   Target: {a['target_by']}" if a.get("target_by") else ""), size=7.5, color=GREY, after=5)
    pas = run.get("prior_actions_status") or []
    if pas:
        kicker(doc, f"Status of actions from {hist[-1]['period'] if hist else 'the prior report'}")
        ts = table(doc, ["Action", "Status", "Evidence and metrics"], [5.6, 2.2, 9.6])
        for pa in pas:
            row = new_row(ts); set_cell(row[0], pa.get("heading", pa.get("id")), bold=True, color=NAVY)
            stt = pa.get("status", ""); col = {"completed": TEAL, "on_track": BLUE, "behind": AMBER, "not_started": "A6192E", "cancelled": GREY, "no_public_evidence": SLATE}.get(stt, GREY)
            set_cell(row[1], stt.replace("_", " "), bold=True, color="FFFFFF", align="center"); shade(row[1], col)
            p = row[2].paragraphs[0]; text_with_links(p, (pa.get("evidence_note") or "") + ((" Metrics: " + pa["metrics_note"]) if pa.get("metrics_note") else ""), size=7.5)
    elif hist and hist[-1].get("actions") and not io:
        kicker(doc, f"Actions from {hist[-1]['period']}"); para(doc, "Status of the prior actions cannot be established from public information; it will be asked for in the inside-out stage.", size=8, italic=True, color=GREY)

    # ================= BACK ON TRACK =================
    bot = run.get("back_on_track") or []
    if bot:
        section(doc, "Getting back on track", "Actions behind or not started: what happened and how to recover")
        tb = table(doc, ["Action", "What happened", "Impact", "Root cause", "Mechanism to prevent", "Measures to track"], [3.0, 3.1, 2.6, 2.9, 3.0, 2.8], hsize=7)
        for b in bot:
            row = new_row(tb); set_cell(row[0], b.get("heading", b.get("id")), bold=True, color=NAVY, size=7.5)
            for k, key in enumerate(("what_happened", "impact", "root_cause", "prevention", "measures"), 1):
                v = b.get(key, ""); set_cell(row[k], "; ".join(v) if isinstance(v, list) else v, size=7.5)

    # ================= ADDITIONAL INFO =================
    ain = run.get("additional_information_needed") or []
    if ain:
        section(doc, "Additional information that would sharpen the assessment", "Only where it would change a grade or an action")
        for b in ain:
            p = doc.add_paragraph(); run_fmt(p.add_run("• " + b), 8.5); p.paragraph_format.space_after = Pt(1)
    doc.add_page_break()

    # ================= APPENDICES =================
    appn = 0
    def app_title(t): 
        nonlocal appn; appn += 1; section(doc, f"Appendix {appn}", t, 13)
    def data_sources_appendix():
        app_title("Data sources: public, inside-out and simulated")
        para(doc, "Every source used in this report, in one line each, with its origin (public record, inside-out board document, or director self-assessment) and whether it was simulated for prototype testing.", size=7.5, italic=True, color=GREY)
        ta = table(doc, ["Source (one line)", "Origin", "Simulated?", "Used for"], [8.0, 2.4, 2.0, 5.0])
        used = {}
        for iid, ind in comp["indicators"].items():
            for l in lenses:
                for e in ((ind.get(l) or {}).get("evidence") or []): used.setdefault(e, set()).add(iid)
        for eid in sorted(run["evidence"], key=lambda x: (x[:2], len(x), x)):
            e = run["evidence"][eid]; row = new_row(ta); set_cell(row[0], f"{eid}  {e.get('document','')}", size=7.5)
            orig = e.get("source_origin") or ("inside_out" if e.get("evidence_class") == "inside_out" else "self_assessment" if e.get("evidence_class") == "self_assessment" else "public")
            set_cell(row[1], orig.replace("_", " "), size=7.5, align="center"); set_cell(row[2], "Yes" if e.get("simulated") else "No", size=7.5, align="center", bold=bool(e.get("simulated")), color=AMBER if e.get("simulated") else None)
            set_cell(row[3], ", ".join(sorted(used.get(eid, []))) or "context", size=7)
    app_title("Detailed indicator view")
    for d in DIMS:
        kicker(doc, f"{d}. {FW['dimensions'][d]}  ·  " + "  ·  ".join(f"{labels[l]} {C['dimensions'][d][l]['grade'] or 'n/a'}" for l in lenses))
        atext = (comp.get("assessment_text") or {}).get(d)
        if atext: para(doc, atext, size=8.5)
        hdr = ["Indicator"] + [labels[l] for l in lenses] + ["Rationale and sources", "Conf."]; widths = [3.2] + [1.0] * len(lenses) + [17.4 - 3.2 - 1.0 * len(lenses) - 1.6, 1.6]
        td = table(doc, hdr, widths)
        for i in [x for x in IND if x["dimension"] == d]:
            ind = comp["indicators"].get(i["id"], {}); row = new_row(td); set_cell(row[0], f"{i['id']} {i['name']}", bold=True, color=NAVY); k = 1
            for l in lenses: score_cell(row[k], (ind.get(l) or {}).get("score"), 9); k += 1
            rc = row[k]; rc.text = ""; first = True
            for l in lenses:
                cl = ind.get(l) or {}; p = rc.paragraphs[0] if first else rc.add_paragraph(); first = False
                lead = f"{labels[l]}: " if len(lenses) > 1 else None
                if cl.get("score") is None: run_fmt(p.add_run((lead or "") + (cl.get("rationale") or "Not assessable.") + (f" [{cl.get('na_cause')}]" if cl.get("na_cause") else "")), 7.5, italic=True, color=GREY)
                else: text_with_links(p, cl.get("rationale") or "", cl.get("evidence") or [], 7.5, lead)
            set_cell(row[k + 1], " / ".join((ind.get(l) or {}).get("confidence") or "–" for l in lenses), size=7.5, align="center")
        lim = (comp.get("limitations") or {}).get(d)
        if lim: para(doc, "Where this grade is limited: " + lim, size=7.5, italic=True, color=GREY)
    # peer one-pagers
    appn += 1
    for n, p_ in enumerate(peers):
        doc.add_page_break(); pc = p_["computed"]
        section(doc, f"Appendix {appn}" + (f" ({n+1} of {len(peers)})"), f"{p_['name']}: outside-in indicator view", 13)
        if io: para(doc, "Carried forward from the outside-in report; peers are not re-assessed with inside-out information.", size=7.5, italic=True, color=GREY)
        gt = doc.add_table(rows=2, cols=7); borders(gt, "none"); fix_widths(gt, [2.48] * 7)
        for k, d in enumerate(DIMS): grade_cell(gt.rows[0].cells[k], pc["dimensions"][d]["board"]["grade"], 12); set_cell(gt.rows[1].cells[k], f"{d}. {FW['dimensions'][d]}", size=6.5, color=GREY, align="center")
        grade_cell(gt.rows[0].cells[6], pc["overall"]["board"].get("grade"), 12); set_cell(gt.rows[1].cells[6], "Overall", size=6.5, color=GREY, align="center", bold=True)
        if p_.get("summary"): para(doc, p_["summary"], size=8.5, before=4)
        tq = table(doc, ["Indicator", "Score", "Rationale and sources"], [4.4, 1.1, 11.9])
        for i in IND:
            cl = (p_["indicators"].get(i["id"]) or {}).get("board") or {}; row = new_row(tq); set_cell(row[0], f"{i['id']} {i['name']}", bold=True, color=NAVY, size=7.5); score_cell(row[1], cl.get("score"), 8)
            p = row[2].paragraphs[0]
            if cl.get("score") is None: run_fmt(p.add_run(cl.get("rationale") or "Not assessable."), 7, italic=True, color=GREY)
            else: text_with_links(p, cl.get("rationale") or "", cl.get("evidence") or [], 7)
    if io:
        doc.add_page_break(); data_sources_appendix()
    # how to read
    doc.add_page_break(); app_title("How to read this report")
    para(doc, "This assessment applies Jeff Bezos's Day 1 / Day 2 lens to a listed company's board, using " + ("public information only." if not io else "public information plus board documents and a director self-assessment supplied by the board.") + " Day 1 boards stay obsessed with the customer, make high-velocity decisions, resist proxies and process, embrace external trends, build rather than protect, and manage to cash. Day 2 is what creeps in afterwards.", size=8.5)
    para(doc, "Six dimensions, 22 indicators. Each indicator is scored 1 (Day 2), 3 (Mixed) or 5 (Day 1) against written anchors, once, for the board. Management's observable actions are attributed to the board's oversight, because a listed board is accountable for what the company does under its watch. Scores of 2 and 4 are used only where evidence meets parts of both adjacent anchors.", size=8.5)
    tg = table(doc, ["Average indicator score", "Grade", "Classification"], [5, 2, 10.4])
    for b in FW["grade_scale"]: row = new_row(tg); set_cell(row[0], FW["grade_scale_ranges"][b["grade"]]); grade_cell(row[1], b["grade"], 9); set_cell(row[2], b["label"])
    para(doc, "Overall grades are the rounded average of the dimension letters (A=5 … F=1; midpoints round down); provisional where only three dimensions are reported. Where evidence falls below the minimum, an indicator is marked not assessable and excluded rather than scored as Mixed, with the cause recorded (disclosure-limited, retrieval-limited, rule-limited). Peers are assessed on public information only.", size=8.5, before=4)
    app_title(f"Methodology and scoring rules (v{FW['rule_version']})")
    rules = ["Every indicator is scored once, for the board; a cell with nothing to assess is marked not applicable.", "Scores are 1, 3 or 5; 2 and 4 only where evidence meets parts of both adjacent anchors.", "A 1 or 5 requires two independent items from two documents, or one strong direct item (regulator or court record, formal board act, audited figure, or a board document); a 3 requires one.", "Confidence is assigned on the evidence: High = two direct sources or one strong source; Medium = one direct or two indirect; Low = a single indirect source, an inference, a silence-based score, or self-assessment alone.", "Below the minimum an indicator is not assessable and excluded, never defaulted to 3; silence is scorable in full on 1.4, 5.2, 5.3 and capped at 3 on 2.3, 2.4, 4.3; an empty register search supports up to 3 on 1.3.", "A dimension grade is reported only if at least half (outside-in) or two-thirds (post inside-out) of its indicators are assessed.", "Overall grade is the rounded average of dimension letters; provisional at three dimensions; none below three.", "Behaviour indicators use 24 months (36 for 2.1 and 4.3); trend indicators five years of filings.", "3.2, 5.1, 6.1 and 6.2 are scored relative to the named peer set.", "Within scoped pairs an evidence item is primary for one indicator only.", "Approved news and trade press (for event facts) are direct evidence; formal board acts are strong direct evidence.", "Inside-out: board documents outrank public sources; the post inside-out score must cite at least one inside-out or self-assessment item; self-assessment alone is at most Medium confidence; peers are never re-scored.", "Each company is run twice on the same rule version; indicator scores within one point and unchanged dimension letters are required for the result to stand."]
    for k, rt in enumerate(rules, 1):
        p = doc.add_paragraph(); run_fmt(p.add_run(f"{k:02d}  "), 8, bold=True, color=BLUE); run_fmt(p.add_run(rt), 8); p.paragraph_format.space_after = Pt(2)
    # sources with bookmarks
    app_title("Sources consulted"); para(doc, "Each entry is referenced from the tables by its code. Verbatim excerpts are kept under 15 words.", size=7.5, italic=True, color=GREY)
    used = {}
    for s in run["subjects"]:
        for iid, ind in (s.get("indicators") or {}).items():
            for l in ("board", "board_io"):
                for e in ((ind.get(l) or {}).get("evidence") or []): used.setdefault(e, set()).add(iid + ("" if s["role"] == "company" else f" ({s['name']})"))
    bid = 100
    for eid in sorted(run["evidence"], key=lambda x: (x[:2], len(x), x)):
        e = run["evidence"][eid]; p = doc.add_paragraph(); add_bookmark(p, f"EV_{eid}", bid); bid += 1
        run_fmt(p.add_run(f"{eid}  "), 8.5, bold=True, color=BLUE); run_fmt(p.add_run(e.get("document", "")), 8.5, bold=True, color=NAVY)
        run_fmt(p.add_run(f"   {e.get('date','')}   ·   {e.get('evidence_class','direct')}" + (f"   ·   {e.get('source_origin')}" if e.get("source_origin") else "") + ("   ·   SIMULATED" if e.get("simulated") else "") + ("   ·   RECENT" if e.get("recent") else "")), 7.5, color=GREY)
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
        if e.get("question"): para(doc, "Q: " + e["question"], size=7.5, italic=True, color=GREY, after=0)
        if e.get("excerpt"): para(doc, e["excerpt"], size=8, italic=True, after=0)
        if e.get("search_log"): para(doc, "Search: " + e["search_log"], size=7.5, color=GREY, after=0)
        p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        if e.get("url"): add_external_link(p2, e["url"], doc.part, 7)
        para(doc, "Used in: " + (", ".join(sorted(used.get(eid, []))) or "context"), size=7, color=GREY, after=5)
    # anchors
    app_title("Indicator definitions and anchors")
    for d in DIMS:
        kicker(doc, f"{d}. {FW['dimensions'][d]}")
        for i in [x for x in IND if x["dimension"] == d]:
            ov = (FW.get("outside_in_overrides") or {}).get(i["id"]); an = ov["anchors"] if (ov and not io) else i["anchors"]
            para(doc, f"{i['id']} {i['name']}", size=8.5, bold=True, color=NAVY, after=1, before=3, keep=True); para(doc, i["definition"] + ("  (Outside-in anchors applied.)" if ov and not io else ""), size=7.5, after=1)
            ta2 = table(doc, ["1 – Day 2", "3 – Mixed", "5 – Day 1"], [5.8, 5.8, 5.8], hsize=7); row = new_row(ta2)
            for k in range(3): set_cell(row[k], an[str([1, 3, 5][k])], size=7)
            para(doc, "Board prompt: " + i.get("board_lens", ""), size=7, color=GREY, after=3)
    doc.save(out); print(f"Report written to {out}")


if __name__ == "__main__":
    main()
