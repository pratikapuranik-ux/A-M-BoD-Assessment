"""Text extraction from uploaded board documents (.docx, .pdf, .xlsx, .txt/.md/.csv)."""
from __future__ import annotations

import base64
import io
import os

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json"}
MIME_EXT = {"application/pdf": ".pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx", "text/plain": ".txt", "text/markdown": ".md", "text/csv": ".csv"}
SUPPORTED = sorted(TEXT_EXT | {".pdf", ".docx", ".xlsx"})
MAX_FILE = 25 * 1024 * 1024  # any one uploaded file
MAX_UPLOAD = 100 * 1024 * 1024  # all documents of one inside-out request together
# Claude reads a PDF as pages only within one request's limits (about 32 MB encoded, 100 pages);
# PDFs beyond this budget are sent as their extracted text instead.
PDF_INLINE_BYTES, PDF_INLINE_PAGES = 18 * 1024 * 1024, 100


class ExtractError(ValueError):
    pass


def _docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(io.BytesIO(data)); out = []
    for el in doc.element.body.iterchildren():  # body order, so tables stay next to their headings
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(el, doc).text.strip()
            if t: out.append(t)
        elif tag == "tbl":
            for row in Table(el, doc).rows:
                cells, seen = [], set()
                for c in row.cells:  # merged cells repeat; keep each once
                    if id(c._tc) in seen: continue
                    seen.add(id(c._tc)); cells.append(c.text.strip().replace("\n", " "))
                if any(cells): out.append(" | ".join(cells))
    return "\n".join(out)


def _pdf(data: bytes, name: str) -> tuple[int, str]:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return len(r.pages), "\n\n".join(f"[page {k}]\n{(p.extract_text() or '').strip()}" for k, p in enumerate(r.pages, 1))


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True); out = []
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(vals): out.append(" | ".join(vals))
    return "\n".join(out)


def _text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try: return data.decode(enc)
        except UnicodeDecodeError: continue
    return data.decode("utf-8", "replace")


def ext_of(name: str, mime: str | None = None) -> str:
    e = os.path.splitext(name or "")[1].lower()
    return e if e else MIME_EXT.get((mime or "").split(";")[0].strip().lower(), "")


def extract(name: str, data: bytes, slot: int = 0, mime: str | None = None) -> dict:
    """-> {name, slot, text, pdf_b64|None}. Raises ExtractError for unsupported or unreadable files."""
    e = ext_of(name, mime); pages = 0
    if len(data) > MAX_FILE: raise ExtractError(f"{name}: {len(data) / 2**20:.0f} MB; the limit is {MAX_FILE // 2**20} MB per file")
    try:
        if e == ".docx": text = _docx(data)
        elif e == ".pdf": pages, text = _pdf(data, name)
        elif e == ".xlsx": text = _xlsx(data)
        elif e in TEXT_EXT: text = _text(data)
        else: raise ExtractError(f"{name}: unsupported file type (supported: {', '.join(SUPPORTED)})")
    except ExtractError: raise
    except Exception as ex:  # noqa: BLE001
        raise ExtractError(f"{name}: could not be read ({type(ex).__name__}: {ex})") from ex
    if e != ".pdf" and not text.strip(): raise ExtractError(f"no readable text in {name}; export it as PDF or text")
    return {"name": name, "slot": int(slot), "text": text.strip(), "pdf_b64": base64.b64encode(data).decode("ascii") if e == ".pdf" else None,
            "pages": pages, "size": len(data)}


def fit_pdfs(docs: list[dict]) -> list[dict]:
    """Keep PDFs as pages while they fit one request, smallest first; send the rest as text. A scanned PDF with no text that does not fit is an error."""
    used_b = used_p = 0
    for d in sorted((d for d in docs if d.get("pdf_b64")), key=lambda d: d["size"]):
        if used_b + d["size"] <= PDF_INLINE_BYTES and used_p + d["pages"] <= PDF_INLINE_PAGES:
            used_b += d["size"]; used_p += d["pages"]; continue
        if not "".join(l for l in d["text"].splitlines() if not l.startswith("[page ")).strip():
            raise ExtractError(f"{d['name']}: a scanned PDF ({d['pages']} pages) with no readable text, too large to read as pages alongside the other PDFs; send it on its own, split it, or send a text version")
        d["pdf_b64"] = None
    return docs


def extract_b64(doc: dict) -> dict:
    """API document {slot, name, mime, data_b64} -> extract()."""
    try: data = base64.b64decode(doc.get("data_b64") or "", validate=False)
    except Exception as ex: raise ExtractError(f"{doc.get('name')}: data_b64 is not valid base64") from ex
    if not data: raise ExtractError(f"{doc.get('name') or 'document'}: empty file")
    return extract(doc.get("name") or "document", data, int(doc.get("slot") or 0), doc.get("mime"))
